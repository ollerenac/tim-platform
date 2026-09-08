"""Generate and validate sealed, independent EXP-02 annotation workbooks."""

from __future__ import annotations

from collections.abc import Mapping
from dataclasses import dataclass
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import re
import stat
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from .jsonio import _open_new_file, canonical_bytes, load_json, sha256_file
from .records import DocumentRecord


SHEET_NAMES = ["LEEME", "DOCUMENTOS", "ENTIDADES", "RELACIONES", "DUDAS", "_LISTAS"]
ENTITY_COLUMNS = [
    "document_id", "entity_id", "entity_type", "indicator_subtype", "normalized_value",
    "mention_as_written", "supporting_quote", "page_or_lines", "certainty", "notes",
]
RELATIONSHIP_COLUMNS = [
    "document_id", "relationship_id", "source_entity_id", "relationship_type",
    "target_entity_id", "supporting_quote", "page_or_lines", "certainty", "notes",
]
DOCUMENT_COLUMNS = [
    "document_id", "source_id", "title", "author", "published_at", "origin_url", "text_path",
    "media_type", "word_count", "minutes", "notes",
]
QUESTION_COLUMNS = ["document_id", "page_or_lines", "doubt", "notes"]
MAX_ANNOTATION_ROWS = 1000
IDENTIFIER = re.compile(r"^[a-z0-9][a-z0-9-]{2,63}$")
ANNOTATION_SHEET_LAYOUTS: dict[str, tuple[dict[str, float], frozenset[int]]] = {
    "DOCUMENTOS": (
        {"A": 22, "B": 18, "C": 40, "D": 24, "E": 22, "F": 48, "G": 62,
         "H": 14, "I": 13, "J": 12, "K": 36},
        frozenset({3, 6, 7, 11}),
    ),
    "ENTIDADES": (
        {"A": 22, "B": 18, "C": 20, "D": 18, "E": 32, "F": 32, "G": 56,
         "H": 18, "I": 14, "J": 36},
        frozenset({5, 6, 7, 8, 10}),
    ),
    "RELACIONES": (
        {"A": 22, "B": 22, "C": 20, "D": 20, "E": 20, "F": 56, "G": 18,
         "H": 14, "I": 36},
        frozenset({6, 7, 9}),
    ),
    "DUDAS": (
        {"A": 22, "B": 18, "C": 60, "D": 42},
        frozenset({2, 3, 4}),
    ),
}


class WorkbookError(ValueError):
    """A workbook cannot be safely imported as an EXP-02 annotation."""


@dataclass(frozen=True)
class WorkbookSummary:
    manifest_digest: str
    annotator_id: str
    document_count: int
    entity_count: int
    relationship_count: int
    question_count: int


@dataclass(frozen=True)
class _Manifest:
    digest: str
    documents: tuple[DocumentRecord, ...]


def _config_path() -> Path:
    return Path(__file__).resolve().parents[2] / "config" / "annotation-types.v1.json"


def _annotation_types() -> dict[str, tuple[str, ...]]:
    try:
        raw = json.loads(_config_path().read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise WorkbookError("annotation vocabulary is unavailable") from error
    required = {"entity_types", "indicator_subtypes", "relationship_types", "certainty_values"}
    if raw.get("schema_version") != 1 or set(raw) != required | {"schema_version"}:
        raise WorkbookError("invalid annotation vocabulary schema")
    result: dict[str, tuple[str, ...]] = {}
    for name in required:
        values = raw[name]
        if not isinstance(values, list) or not values or any(not isinstance(value, str) for value in values):
            raise WorkbookError(f"invalid annotation vocabulary: {name}")
        if len(values) != len(set(values)):
            raise WorkbookError(f"duplicate annotation vocabulary value: {name}")
        result[name] = tuple(values)
    return result


def _load_manifest(manifest: Mapping[str, object] | str | Path) -> _Manifest:
    if isinstance(manifest, (str, Path)):
        raw = load_json(manifest)
        digest = sha256_file(manifest)
    elif isinstance(manifest, Mapping):
        raw = dict(manifest)
        digest = hashlib.sha256(canonical_bytes(raw)).hexdigest()
    else:
        raise WorkbookError("manifest must be a mapping or frozen JSON path")
    if raw.get("schema_version") != 1 or not isinstance(raw.get("documents"), list):
        raise WorkbookError("invalid input manifest")
    try:
        documents = tuple(DocumentRecord.from_dict(value) for value in raw["documents"])
    except (TypeError, ValueError) as error:
        raise WorkbookError(f"invalid input manifest document: {error}") from error
    if len(documents) != 16:
        raise WorkbookError("input manifest must contain exactly 16 documents")
    document_ids = [document.document_id for document in documents]
    if len(document_ids) != len(set(document_ids)):
        raise WorkbookError("input manifest has duplicate document_id")
    if any(document.role != "primary" for document in documents):
        raise WorkbookError("input manifest documents must all be primary")
    return _Manifest(digest, documents)


def _safe_save_new(workbook: Workbook, output: Path) -> None:
    buffer = BytesIO()
    workbook.save(buffer)
    try:
        descriptor = _open_new_file(output)
        with os.fdopen(descriptor, "xb") as handle:
            handle.write(buffer.getvalue())
    except FileExistsError as error:
        raise WorkbookError("refusing to overwrite annotation workbook") from error
    except ValueError as error:
        raise WorkbookError(str(error)) from error
    except OSError as error:
        raise WorkbookError(f"cannot write workbook: {output}") from error


def _style_header(sheet, columns: list[str]) -> None:  # type: ignore[no-untyped-def]
    fill = PatternFill("solid", fgColor="1F4E78")
    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(1, index, name)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.protection = Protection(locked=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{chr(64 + len(columns))}{MAX_ANNOTATION_ROWS + 1}"
    sheet.protection.sheet = True


def _set_widths(sheet, widths: dict[str, float]) -> None:  # type: ignore[no-untyped-def]
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _unlock_cells(sheet, columns: int) -> None:  # type: ignore[no-untyped-def]
    for row in range(2, MAX_ANNOTATION_ROWS + 2):
        for column in range(1, columns + 1):
            sheet.cell(row, column).protection = Protection(locked=False)


def configure_annotation_sheet(sheet, columns: list[str], *, unlock: bool) -> None:  # type: ignore[no-untyped-def]
    """Apply the one canonical, LibreOffice-usable annotation table layout."""
    try:
        widths, wrapped_columns = ANNOTATION_SHEET_LAYOUTS[sheet.title]
    except KeyError as error:
        raise WorkbookError(f"unknown annotation sheet layout: {sheet.title}") from error
    _style_header(sheet, columns)
    _set_widths(sheet, widths)
    sheet.row_dimensions[1].height = 30
    sheet.sheet_format.defaultRowHeight = 32
    if not unlock:
        return
    _unlock_cells(sheet, len(columns))
    for row in range(2, MAX_ANNOTATION_ROWS + 2):
        for column in wrapped_columns:
            sheet.cell(row, column).alignment = Alignment(wrap_text=True, vertical="top")


def _add_list_validation(sheet, column: str, formula: str) -> None:  # type: ignore[no-untyped-def]
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "Select a value from the closed EXP-02 vocabulary."
    validation.errorTitle = "Invalid annotation value"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"{column}2:{column}{MAX_ANNOTATION_ROWS + 1}")


def _build_lists(sheet, vocabulary: Mapping[str, tuple[str, ...]], documents: tuple[DocumentRecord, ...]) -> None:  # type: ignore[no-untyped-def]
    columns = [
        ("entity_types", vocabulary["entity_types"]),
        ("relationship_types", vocabulary["relationship_types"]),
        ("certainty_values", vocabulary["certainty_values"]),
        ("indicator_subtypes", vocabulary["indicator_subtypes"]),
        ("primary_document_ids", tuple(document.document_id for document in documents)),
    ]
    for index, (name, values) in enumerate(columns, start=1):
        sheet.cell(1, index, name).protection = Protection(locked=True)
        for row, value in enumerate(values, start=2):
            sheet.cell(row, index, value).protection = Protection(locked=True)
    sheet.sheet_state = "hidden"
    sheet.protection.sheet = True


def build_workbook(
    manifest: Mapping[str, object] | str | Path,
    annotator_id: str,
    output_path: str | Path,
) -> str:
    """Create one independent annotation workbook bound to a frozen manifest."""
    frozen = _load_manifest(manifest)
    if annotator_id not in {"annotator-a", "annotator-b"}:
        raise WorkbookError("annotator_id must be annotator-a or annotator-b")
    output = Path(output_path)
    if output.suffix.lower() != ".xlsx" or output.stem != annotator_id:
        raise WorkbookError("output filename must be the annotator identity plus .xlsx")
    vocabulary = _annotation_types()
    documents = frozen.documents

    workbook = Workbook()
    readme = workbook.active
    readme.title = "LEEME"
    document_sheet = workbook.create_sheet("DOCUMENTOS")
    entity_sheet = workbook.create_sheet("ENTIDADES")
    relationship_sheet = workbook.create_sheet("RELACIONES")
    doubts_sheet = workbook.create_sheet("DUDAS")
    lists_sheet = workbook.create_sheet("_LISTAS")

    readme["A1"] = "EXP-02 annotation workbook"
    readme["A2"] = "annotator_id"
    readme["B2"] = annotator_id
    readme["A3"] = "input_manifest_sha256"
    readme["B3"] = frozen.digest
    readme["A4"] = "document_scope"
    readme["B4"] = "fixed-16-primary"
    instructions = [
        "Trabaje de forma independiente: no comparta esta hoja ni discuta casos durante la primera pasada.",
        "Marque solo afirmaciones explícitas del documento; conserve una frase exacta y la página o líneas.",
        "Registre los minutos activos una sola vez por documento en DOCUMENTOS.",
        "Use DUDAS en vez de adivinar; la incertidumbre no elimina una fila anotada.",
    ]
    for row, text in enumerate(instructions, start=5):
        readme.cell(row, 1, text).alignment = Alignment(wrap_text=True, vertical="top")
    _set_widths(readme, {"A": 34, "B": 90})
    readme.protection.sheet = True
    for row in readme.iter_rows():
        for cell in row:
            cell.protection = Protection(locked=True)

    configure_annotation_sheet(document_sheet, DOCUMENT_COLUMNS, unlock=False)
    for row, document in enumerate(documents, start=2):
        values = [
            document.document_id, document.source_id, document.title, document.author,
            document.published_at, document.origin_url, document.text_path, document.media_type,
            document.word_count, None, None,
        ]
        for column, value in enumerate(values, start=1):
            cell = document_sheet.cell(row, column, value)
            cell.alignment = Alignment(wrap_text=column in {3, 6, 7, 11}, vertical="top")
            cell.protection = Protection(locked=column not in {10, 11})

    configure_annotation_sheet(entity_sheet, ENTITY_COLUMNS, unlock=True)

    configure_annotation_sheet(relationship_sheet, RELATIONSHIP_COLUMNS, unlock=True)

    configure_annotation_sheet(doubts_sheet, QUESTION_COLUMNS, unlock=True)

    _build_lists(lists_sheet, vocabulary, documents)
    document_ids_formula = f"'_LISTAS'!$E$2:$E${len(documents) + 1}"
    _add_list_validation(entity_sheet, "A", document_ids_formula)
    _add_list_validation(entity_sheet, "C", "'_LISTAS'!$A$2:$A$12")
    _add_list_validation(entity_sheet, "D", "'_LISTAS'!$D$2:$D$8")
    _add_list_validation(entity_sheet, "I", "'_LISTAS'!$C$2:$C$3")
    _add_list_validation(relationship_sheet, "A", document_ids_formula)
    _add_list_validation(relationship_sheet, "D", "'_LISTAS'!$B$2:$B$6")
    _add_list_validation(relationship_sheet, "H", "'_LISTAS'!$C$2:$C$3")
    _add_list_validation(doubts_sheet, "A", document_ids_formula)

    _safe_save_new(workbook, output)
    return str(output)


def _require_regular_xlsx(path: str | Path) -> Path:
    output = Path(path)
    try:
        status = os.lstat(output)
    except OSError as error:
        raise WorkbookError("annotation workbook is unavailable") from error
    if stat.S_ISLNK(status.st_mode) or not stat.S_ISREG(status.st_mode):
        raise WorkbookError("annotation workbook must be a regular non-symlink file")
    if output.suffix.lower() != ".xlsx":
        raise WorkbookError("annotation workbook must have .xlsx extension")
    return output


def _reject_unsafe_archive(path: Path) -> None:
    try:
        with ZipFile(path) as archive:
            names = archive.namelist()
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise WorkbookError("annotation workbook contains macros")
            if any(name.startswith("xl/externalLinks/") for name in names):
                raise WorkbookError("annotation workbook contains external links")
            for name in names:
                if not name.endswith(".rels"):
                    continue
                try:
                    relationships = ElementTree.fromstring(archive.read(name))
                except ElementTree.ParseError as error:
                    raise WorkbookError("annotation workbook has invalid relationship metadata") from error
                if any(relationship.attrib.get("TargetMode") == "External" for relationship in relationships):
                    raise WorkbookError("annotation workbook contains external links")
    except (OSError, BadZipFile) as error:
        raise WorkbookError("annotation workbook is not a valid .xlsx archive") from error


def _load_workbook(path: Path):  # type: ignore[no-untyped-def]
    _reject_unsafe_archive(path)
    try:
        workbook = load_workbook(path, data_only=False, keep_links=True)
    except (OSError, BadZipFile, ValueError) as error:
        raise WorkbookError("annotation workbook cannot be read") from error
    if workbook._external_links:
        raise WorkbookError("annotation workbook contains external links")
    return workbook


def _values(sheet, row: int, length: int) -> list[object]:  # type: ignore[no-untyped-def]
    return [sheet.cell(row, column).value for column in range(1, length + 1)]


def _validate_header(sheet, expected: list[str]) -> None:  # type: ignore[no-untyped-def]
    if _values(sheet, 1, len(expected)) != expected:
        raise WorkbookError(f"invalid {sheet.title} columns")


def _nonempty(value: object, field: str, sheet: str, row: int) -> str:
    if not isinstance(value, str) or not value.strip():
        raise WorkbookError(f"{sheet} row {row}: {field} is required")
    return value


def _row_has_value(values: list[object]) -> bool:
    return any(value not in (None, "") for value in values)


def _assert_no_formulas(workbook) -> None:  # type: ignore[no-untyped-def]
    for sheet_name in ("DOCUMENTOS", "ENTIDADES", "RELACIONES", "DUDAS"):
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    raise WorkbookError(f"formula in user cell: {sheet_name}!{cell.coordinate}")


def _validate_documents(sheet, primary: tuple[DocumentRecord, ...]) -> None:  # type: ignore[no-untyped-def]
    _validate_header(sheet, DOCUMENT_COLUMNS)
    expected_rows = [
        [document.document_id, document.source_id, document.title, document.author, document.published_at,
         document.origin_url, document.text_path, document.media_type, document.word_count]
        for document in primary
    ]
    for index, expected in enumerate(expected_rows, start=2):
        actual = _values(sheet, index, 9)
        if actual != expected:
            raise WorkbookError("DOCUMENTOS identity does not match input manifest")
        minutes = sheet.cell(index, 10).value
        if minutes not in (None, "") and (isinstance(minutes, bool) or not isinstance(minutes, int) or minutes < 0):
            raise WorkbookError(f"DOCUMENTOS row {index}: minutes must be a non-negative integer")
    for index in range(len(expected_rows) + 2, sheet.max_row + 1):
        if _row_has_value(_values(sheet, index, len(DOCUMENT_COLUMNS))):
            raise WorkbookError("DOCUMENTOS contains an unexpected document")


def _validate_entities(sheet, document_ids: set[str], vocabulary: Mapping[str, tuple[str, ...]]) -> dict[tuple[str, str], int]:  # type: ignore[no-untyped-def]
    _validate_header(sheet, ENTITY_COLUMNS)
    entities: dict[tuple[str, str], int] = {}
    for row in range(2, sheet.max_row + 1):
        values = _values(sheet, row, len(ENTITY_COLUMNS))
        if not _row_has_value(values):
            continue
        document_id = _nonempty(values[0], "document_id", "ENTIDADES", row)
        if document_id not in document_ids:
            raise WorkbookError(f"ENTIDADES row {row}: unknown document_id")
        entity_id = _nonempty(values[1], "entity_id", "ENTIDADES", row)
        key = (document_id, entity_id)
        if key in entities:
            raise WorkbookError(f"ENTIDADES row {row}: duplicate entity_id")
        entities[key] = row
        entity_type = _nonempty(values[2], "entity_type", "ENTIDADES", row)
        if entity_type not in vocabulary["entity_types"]:
            raise WorkbookError(f"ENTIDADES row {row}: invalid entity_type")
        subtype = values[3]
        if entity_type == "indicator":
            if subtype not in vocabulary["indicator_subtypes"]:
                raise WorkbookError(f"ENTIDADES row {row}: invalid indicator_subtype")
        elif subtype not in (None, ""):
            raise WorkbookError(f"ENTIDADES row {row}: indicator_subtype only applies to indicators")
        _nonempty(values[4], "normalized_value", "ENTIDADES", row)
        _nonempty(values[5], "mention_as_written", "ENTIDADES", row)
        _nonempty(values[6], "supporting_quote", "ENTIDADES", row)
        _nonempty(values[7], "page_or_lines", "ENTIDADES", row)
        certainty = _nonempty(values[8], "certainty", "ENTIDADES", row)
        if certainty not in vocabulary["certainty_values"]:
            raise WorkbookError(f"ENTIDADES row {row}: invalid certainty")
    return entities


def _validate_relationships(sheet, document_ids: set[str], entities: dict[tuple[str, str], int], vocabulary: Mapping[str, tuple[str, ...]]) -> int:  # type: ignore[no-untyped-def]
    _validate_header(sheet, RELATIONSHIP_COLUMNS)
    relationship_ids: set[tuple[str, str]] = set()
    count = 0
    for row in range(2, sheet.max_row + 1):
        values = _values(sheet, row, len(RELATIONSHIP_COLUMNS))
        if not _row_has_value(values):
            continue
        document_id = _nonempty(values[0], "document_id", "RELACIONES", row)
        if document_id not in document_ids:
            raise WorkbookError(f"RELACIONES row {row}: unknown document_id")
        relationship_id = _nonempty(values[1], "relationship_id", "RELACIONES", row)
        key = (document_id, relationship_id)
        if key in relationship_ids:
            raise WorkbookError(f"RELACIONES row {row}: duplicate relationship_id")
        relationship_ids.add(key)
        source = _nonempty(values[2], "source_entity_id", "RELACIONES", row)
        target = _nonempty(values[4], "target_entity_id", "RELACIONES", row)
        if (document_id, source) not in entities or (document_id, target) not in entities:
            raise WorkbookError(f"RELACIONES row {row}: endpoint not found in document")
        relationship_type = _nonempty(values[3], "relationship_type", "RELACIONES", row)
        if relationship_type not in vocabulary["relationship_types"]:
            raise WorkbookError(f"RELACIONES row {row}: invalid relationship_type")
        _nonempty(values[5], "supporting_quote", "RELACIONES", row)
        _nonempty(values[6], "page_or_lines", "RELACIONES", row)
        certainty = _nonempty(values[7], "certainty", "RELACIONES", row)
        if certainty not in vocabulary["certainty_values"]:
            raise WorkbookError(f"RELACIONES row {row}: invalid certainty")
        count += 1
    return count


def _validate_doubts(sheet, document_ids: set[str]) -> int:  # type: ignore[no-untyped-def]
    _validate_header(sheet, QUESTION_COLUMNS)
    count = 0
    for row in range(2, sheet.max_row + 1):
        values = _values(sheet, row, len(QUESTION_COLUMNS))
        if not _row_has_value(values):
            continue
        document_id = _nonempty(values[0], "document_id", "DUDAS", row)
        if document_id not in document_ids:
            raise WorkbookError(f"DUDAS row {row}: unknown document_id")
        _nonempty(values[1], "page_or_lines", "DUDAS", row)
        _nonempty(values[2], "doubt", "DUDAS", row)
        count += 1
    return count


def validate_workbook(path: str | Path, manifest: Mapping[str, object] | str | Path) -> WorkbookSummary:
    """Reject malformed edits and return a small import-ready workbook summary."""
    frozen = _load_manifest(manifest)
    vocabulary = _annotation_types()
    output = _require_regular_xlsx(path)
    workbook = _load_workbook(output)
    if workbook.sheetnames != SHEET_NAMES:
        raise WorkbookError("workbook sheets do not match the fixed EXP-02 layout")
    if workbook["_LISTAS"].sheet_state != "hidden":
        raise WorkbookError("annotation vocabulary sheet must be hidden")
    readme = workbook["LEEME"]
    annotator_id = readme["B2"].value
    scope = readme["B4"].value
    if annotator_id not in {"annotator-a", "annotator-b"}:
        raise WorkbookError("workbook annotator identity does not match its filename")
    if scope != "fixed-16-primary":
        raise WorkbookError("workbook document scope is invalid")
    if output.stem != annotator_id:
        raise WorkbookError("workbook annotator identity does not match its filename")
    if readme["B3"].value != frozen.digest:
        raise WorkbookError("workbook manifest digest does not match frozen input manifest")
    if not readme.protection.sheet or not readme["B2"].protection.locked or not readme["B3"].protection.locked or not readme["B4"].protection.locked:
        raise WorkbookError("workbook identity cells must be protected")
    _assert_no_formulas(workbook)
    documents = frozen.documents
    _validate_documents(workbook["DOCUMENTOS"], documents)
    document_ids = {document.document_id for document in documents}
    entities = _validate_entities(workbook["ENTIDADES"], document_ids, vocabulary)
    relationships = _validate_relationships(workbook["RELACIONES"], document_ids, entities, vocabulary)
    doubts = _validate_doubts(workbook["DUDAS"], document_ids)
    return WorkbookSummary(
        manifest_digest=frozen.digest,
        annotator_id=annotator_id,
        document_count=len(documents),
        entity_count=len(entities),
        relationship_count=relationships,
        question_count=doubts,
    )
