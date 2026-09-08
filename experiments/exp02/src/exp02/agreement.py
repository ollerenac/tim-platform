"""Deterministic, one-to-one agreement calculations for EXP-02 annotations."""

from __future__ import annotations

from collections.abc import Mapping
from dataclasses import asdict, dataclass, replace
from datetime import datetime, timezone
import hashlib
import ipaddress
from pathlib import Path
import re
import unicodedata
from urllib.parse import urlsplit, urlunsplit

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from .annotations import (
    AnnotationSet,
    EntityAnnotation,
    RelationshipAnnotation,
    _snapshot_workbook,
)
from .jsonio import canonical_bytes, write_new_json
from .workbook import (
    WorkbookError,
    _annotation_types,
    _load_workbook,
    _require_regular_xlsx,
    _safe_save_new,
)


def _fold_whitespace(value: str) -> str:
    return " ".join(unicodedata.normalize("NFC", value).split())


def _normalize_url(value: str) -> str:
    parsed = urlsplit(value)
    if not parsed.scheme or not parsed.hostname:
        return value
    credentials = ""
    if parsed.username is not None:
        credentials = parsed.username
        if parsed.password is not None:
            credentials += f":{parsed.password}"
        credentials += "@"
    host = parsed.hostname.casefold()
    if ":" in host and not host.startswith("["):
        host = f"[{host}]"
    try:
        port = f":{parsed.port}" if parsed.port is not None else ""
    except ValueError:
        return value
    return urlunsplit(
        (parsed.scheme.casefold(), f"{credentials}{host}{port}", parsed.path, parsed.query, parsed.fragment)
    )


def normalize_value(entity_type: str, value: str, indicator_subtype: str | None = None) -> str:
    """Produce an explicit comparison key while leaving stored values untouched."""
    value = _fold_whitespace(value)
    if entity_type == "indicator":
        if indicator_subtype == "url":
            return _normalize_url(value)
        if indicator_subtype == "domain":
            return value.rstrip(".").casefold()
        if indicator_subtype == "ip":
            try:
                return ipaddress.ip_address(value).compressed
            except ValueError:
                return value
        if indicator_subtype in {"hash_md5", "hash_sha1", "hash_sha256"}:
            return value.casefold()
        if indicator_subtype == "email":
            local, separator, domain = value.rpartition("@")
            return f"{local}@{domain.casefold()}" if separator and local and domain else value
        return value
    return value.casefold()


def entity_key(entity: EntityAnnotation) -> tuple[str, str, str | None, str]:
    return (
        entity.document_id,
        entity.entity_type,
        entity.indicator_subtype,
        normalize_value(entity.entity_type, entity.normalized_value, entity.indicator_subtype),
    )


def relationship_key(
    relationship: RelationshipAnnotation, entities_by_id: dict[tuple[str, str], EntityAnnotation]
) -> tuple[tuple[str, str, str | None, str], str, tuple[str, str, str | None, str]]:
    return (
        entity_key(entities_by_id[(relationship.document_id, relationship.source_entity_id)]),
        relationship.relationship_type,
        entity_key(entities_by_id[(relationship.document_id, relationship.target_entity_id)]),
    )


@dataclass(frozen=True)
class Metrics:
    tp: int
    fp: int
    fn: int
    precision: float
    recall: float
    f1: float

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


@dataclass(frozen=True)
class QuoteAgreement:
    matched_claims: int
    exact_matches: int
    rate: float

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


@dataclass(frozen=True)
class Disagreement:
    disagreement_id: str
    document_id: str
    claim_kind: str
    annotator_a_value: dict[str, object] | None
    annotator_b_value: dict[str, object] | None

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


@dataclass(frozen=True)
class Comparison:
    annotator_a: AnnotationSet
    annotator_b: AnnotationSet
    entity_metrics: Metrics
    entity_metrics_reversed: Metrics
    relationship_metrics: Metrics
    relationship_metrics_reversed: Metrics
    quote_agreement: QuoteAgreement
    disagreements: tuple[Disagreement, ...]

    def to_dict(self) -> dict[str, object]:
        return {
            "schema_version": 1,
            "annotator_a": self.annotator_a.annotator_id,
            "annotator_b": self.annotator_b.annotator_id,
            "entity_metrics": self.entity_metrics.to_dict(),
            "entity_metrics_reversed": self.entity_metrics_reversed.to_dict(),
            "relationship_metrics": self.relationship_metrics.to_dict(),
            "relationship_metrics_reversed": self.relationship_metrics_reversed.to_dict(),
            "quote_agreement": self.quote_agreement.to_dict(),
            "disagreements": [item.to_dict() for item in self.disagreements],
        }

    @property
    def digest(self) -> str:
        return hashlib.sha256(canonical_bytes(self.to_dict())).hexdigest()


@dataclass(frozen=True)
class ReferenceSet:
    """The adjudicated reference and all provenance needed to seal it."""

    comparison: Comparison
    decisions: tuple[dict[str, object], ...]
    entities: tuple[EntityAnnotation, ...]
    relationships: tuple[RelationshipAnnotation, ...]

    def to_dict(self) -> dict[str, object]:
        return {
            "schema_version": 1,
            "input_manifest_digest": self.comparison.annotator_a.manifest_digest,
            "entities": [item.to_dict() for item in self.entities],
            "relationships": [item.to_dict() for item in self.relationships],
        }


@dataclass(frozen=True)
class ReferenceFreeze:
    output_root: str
    document_count: int
    freeze_digest: str
    frozen_at: str
    artifact_hashes: dict[str, str]

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


def _claim_dict(claim: EntityAnnotation | RelationshipAnnotation) -> dict[str, object]:
    return claim.to_dict()


def _sort_claims(claims: list[EntityAnnotation | RelationshipAnnotation]) -> list[EntityAnnotation | RelationshipAnnotation]:
    return sorted(claims, key=lambda claim: canonical_bytes(_claim_dict(claim)))


def _group_entities(annotations: AnnotationSet) -> dict[tuple[str, str, str | None, str], list[EntityAnnotation]]:
    result: dict[tuple[str, str, str | None, str], list[EntityAnnotation]] = {}
    for entity in annotations.entities:
        result.setdefault(entity_key(entity), []).append(entity)
    return {key: _sort_claims(value) for key, value in result.items()}  # type: ignore[return-value]


def _group_relationships(annotations: AnnotationSet) -> dict[tuple[tuple[str, str, str | None, str], str, tuple[str, str, str | None, str]], list[RelationshipAnnotation]]:
    entities = {(entity.document_id, entity.entity_id): entity for entity in annotations.entities}
    result: dict[tuple[tuple[str, str, str | None, str], str, tuple[str, str, str | None, str]], list[RelationshipAnnotation]] = {}
    for relationship in annotations.relationships:
        result.setdefault(relationship_key(relationship, entities), []).append(relationship)
    return {key: _sort_claims(value) for key, value in result.items()}  # type: ignore[return-value]


def _metric(tp: int, predictions: int, references: int) -> Metrics:
    fp, fn = predictions - tp, references - tp
    precision = tp / predictions if predictions else 0.0
    recall = tp / references if references else 0.0
    f1 = (2 * precision * recall / (precision + recall)) if precision + recall else 0.0
    return Metrics(tp, fp, fn, precision, recall, f1)


def _disagreement(
    claim_kind: str,
    a: EntityAnnotation | RelationshipAnnotation | None,
    b: EntityAnnotation | RelationshipAnnotation | None,
) -> Disagreement:
    document_id = (a or b).document_id  # type: ignore[union-attr]
    a_value = _claim_dict(a) if a else None
    b_value = _claim_dict(b) if b else None
    data = {"claim_kind": claim_kind, "document_id": document_id, "annotator_a_value": a_value, "annotator_b_value": b_value}
    return Disagreement(
        disagreement_id=hashlib.sha256(canonical_bytes(data)).hexdigest(),
        document_id=document_id,
        claim_kind=claim_kind,
        annotator_a_value=a_value,
        annotator_b_value=b_value,
    )


def _compare_groups(
    kind: str,
    a_groups: dict[object, list[object]],
    b_groups: dict[object, list[object]],
) -> tuple[Metrics, Metrics, list[tuple[object, object]], list[Disagreement]]:
    pairs: list[tuple[object, object]] = []
    disagreements: list[Disagreement] = []
    tp = 0
    for key in sorted(set(a_groups) | set(b_groups), key=lambda value: canonical_bytes(value)):
        a_claims, b_claims = a_groups.get(key, []), b_groups.get(key, [])
        for a, b in zip(a_claims, b_claims):
            pairs.append((a, b))
            tp += 1
            if _fold_whitespace(a.supporting_quote) != _fold_whitespace(b.supporting_quote):
                disagreements.append(_disagreement(kind, a, b))
        for a in a_claims[len(b_claims):]:
            disagreements.append(_disagreement(kind, a, None))
        for b in b_claims[len(a_claims):]:
            disagreements.append(_disagreement(kind, None, b))
    return _metric(tp, sum(map(len, a_groups.values())), sum(map(len, b_groups.values()))), _metric(tp, sum(map(len, b_groups.values())), sum(map(len, a_groups.values()))), pairs, disagreements


def compare_annotation_sets(annotator_a: AnnotationSet, annotator_b: AnnotationSet) -> Comparison:
    """Compare independently imported workbooks with stable, one-to-one matching."""
    by_identity = {
        annotator_a.annotator_id: annotator_a,
        annotator_b.annotator_id: annotator_b,
    }
    if set(by_identity) != {"annotator-a", "annotator-b"} or len(by_identity) != 2:
        raise ValueError("annotation sets must contain annotator-a and annotator-b exactly once")
    annotator_a, annotator_b = by_identity["annotator-a"], by_identity["annotator-b"]
    if annotator_a.manifest_digest != annotator_b.manifest_digest:
        raise ValueError("annotation sets must bind the same input manifest")
    if (
        len(annotator_a.document_ids) != 16
        or len(set(annotator_a.document_ids)) != 16
        or len(annotator_b.document_ids) != 16
        or len(set(annotator_b.document_ids)) != 16
    ):
        raise ValueError("annotation sets must each cover exactly 16 documents")
    if annotator_a.document_ids != annotator_b.document_ids:
        raise ValueError("annotation sets must cover the same documents in frozen order")
    entity_metrics, entity_reverse, entity_pairs, entity_disagreements = _compare_groups("entity", _group_entities(annotator_a), _group_entities(annotator_b))
    relationship_metrics, relationship_reverse, relationship_pairs, relationship_disagreements = _compare_groups("relationship", _group_relationships(annotator_a), _group_relationships(annotator_b))
    pairs = entity_pairs + relationship_pairs
    exact = sum(_fold_whitespace(a.supporting_quote) == _fold_whitespace(b.supporting_quote) for a, b in pairs)
    quote = QuoteAgreement(len(pairs), exact, exact / len(pairs) if pairs else 0.0)
    disagreements = tuple(sorted(entity_disagreements + relationship_disagreements, key=lambda item: item.disagreement_id))
    return Comparison(annotator_a, annotator_b, entity_metrics, entity_reverse, relationship_metrics, relationship_reverse, quote, disagreements)


_ADJUDICATION_COLUMNS = [
    "disagreement_id", "document_id", "claim_kind", "annotator_a_value",
    "annotator_b_value", "decision", "final_type", "final_value", "final_quote",
    "page_or_lines", "rationale",
]
_DECISIONS = {"accept-a", "accept-b", "accept-both", "replace", "exclude"}
_DECISION_FIELDS = {"decision", "final_type", "final_value", "final_quote", "page_or_lines", "rationale"}


def build_adjudication_workbook(comparison: Comparison, path: str | Path) -> str:
    """Create a blind decision sheet containing only human-claim differences."""
    output = Path(path)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("adjudication workbook must have .xlsx extension")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ADJUDICACION"
    sheet.append(_ADJUDICATION_COLUMNS)
    validation = DataValidation(type="list", formula1='"accept-a,accept-b,accept-both,replace,exclude"', allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"F2:F{max(2, len(comparison.disagreements) + 1)}")
    for item in comparison.disagreements:
        sheet.append([
            item.disagreement_id, item.document_id, item.claim_kind,
            _display_claim(item.annotator_a_value), _display_claim(item.annotator_b_value),
            None, None, None, None, None, None,
        ])
    try:
        _safe_save_new(workbook, output)
    except WorkbookError as error:
        message = str(error).replace("annotation workbook", "adjudication workbook")
        raise ValueError(message) from error
    return str(output)


def validate_adjudication_workbook(path: str | Path, comparison: Comparison) -> dict[str, dict[str, object]]:
    """Read decisions from the fixed blind workbook and enforce its contract."""
    try:
        snapshot = _snapshot_workbook(path)
    except ValueError as error:
        raise ValueError("adjudication workbook is unavailable") from error
    try:
        return _validate_adjudication_snapshot(snapshot, comparison)
    finally:
        snapshot.unlink(missing_ok=True)
        snapshot.parent.rmdir()


def _validate_adjudication_snapshot(
    path: Path, comparison: Comparison
) -> dict[str, dict[str, object]]:
    try:
        workbook = _load_workbook(_require_regular_xlsx(path))
    except WorkbookError as error:
        raise ValueError(str(error).replace("annotation workbook", "adjudication workbook")) from error
    except (OSError, ValueError) as error:
        raise ValueError("adjudication workbook cannot be read") from error
    if workbook.sheetnames != ["ADJUDICACION"]:
        raise ValueError("adjudication workbook must contain only ADJUDICACION")
    sheet = workbook.active
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f" or (
                isinstance(cell.value, str) and cell.value.startswith("=")
            ):
                raise ValueError(f"formula in adjudication cell: {cell.coordinate}")
    if sheet.max_column != len(_ADJUDICATION_COLUMNS):
        raise ValueError("adjudication workbook must contain exactly the prescribed columns")
    headers = [sheet.cell(1, index).value for index in range(1, len(_ADJUDICATION_COLUMNS) + 1)]
    if headers != _ADJUDICATION_COLUMNS:
        raise ValueError("adjudication workbook columns do not match the fixed layout")
    expected_rows = list(comparison.disagreements)
    for row, item in enumerate(expected_rows, start=2):
        immutable = [item.disagreement_id, item.document_id, item.claim_kind, _display_claim(item.annotator_a_value), _display_claim(item.annotator_b_value)]
        actual = [sheet.cell(row, column).value for column in range(1, 6)]
        if actual != immutable:
            raise ValueError("adjudication workbook has altered immutable claim displays")
    for row in range(len(expected_rows) + 2, sheet.max_row + 1):
        if any(sheet.cell(row, column).value not in (None, "") for column in range(1, len(_ADJUDICATION_COLUMNS) + 1)):
            raise ValueError("adjudication workbook contains non-empty cells outside the prescribed matrix")
    rows: dict[str, dict[str, object]] = {}
    for row in range(2, len(expected_rows) + 2):
        values = {name: sheet.cell(row, index).value for index, name in enumerate(_ADJUDICATION_COLUMNS, start=1)}
        if all(value in (None, "") for value in values.values()):
            continue
        identifier = values["disagreement_id"]
        if not isinstance(identifier, str) or identifier in rows:
            raise ValueError("adjudication workbook has invalid disagreement_id")
        rows[identifier] = {
            name: value
            for name, value in values.items()
            if name not in {"disagreement_id", "annotator_a_value", "annotator_b_value", "document_id", "claim_kind"}
            and value not in (None, "")
        }
    _validate_decision_rows(comparison, rows)
    return rows


def _display_claim(value: dict[str, object] | None) -> str | None:
    if value is None:
        return None
    # Canonical JSON makes a review value reproducible, but deliberately has no
    # experimental condition, run, result, or metric field.
    return canonical_bytes(value).decode("utf-8").rstrip("\n")


def _claim_maps(comparison: Comparison) -> tuple[dict[bytes, EntityAnnotation | RelationshipAnnotation], dict[bytes, EntityAnnotation | RelationshipAnnotation]]:
    a: dict[bytes, EntityAnnotation | RelationshipAnnotation] = {}
    b: dict[bytes, EntityAnnotation | RelationshipAnnotation] = {}
    for annotation, target in ((comparison.annotator_a, a), (comparison.annotator_b, b)):
        for claim in (*annotation.entities, *annotation.relationships):
            target[canonical_bytes(claim.to_dict())] = claim
    return a, b


def _decision_rows(decisions: Mapping[str, object] | list[Mapping[str, object]]) -> dict[str, dict[str, object]]:
    if isinstance(decisions, Mapping):
        result: dict[str, dict[str, object]] = {}
        for identifier, value in decisions.items():
            if not isinstance(identifier, str) or not isinstance(value, Mapping):
                raise ValueError("adjudication decisions must map IDs to decision records")
            result[identifier] = dict(value)
        return result
    if isinstance(decisions, list):
        result = {}
        for value in decisions:
            if not isinstance(value, Mapping) or not isinstance(value.get("disagreement_id"), str):
                raise ValueError("adjudication decision lacks disagreement_id")
            record = dict(value)
            identifier = str(record.pop("disagreement_id"))
            result[identifier] = record
        return result
    raise ValueError("adjudication decisions must be a mapping or list")


def _validate_decision_rows(comparison: Comparison, rows: Mapping[str, Mapping[str, object]]) -> None:
    required_ids = {item.disagreement_id for item in comparison.disagreements}
    if set(rows) != required_ids:
        raise ValueError("adjudication decisions must cover every disagreement exactly once")
    by_id = {item.disagreement_id: item for item in comparison.disagreements}
    vocabulary = _annotation_types()
    for identifier, record in rows.items():
        item = by_id[identifier]
        unknown = set(record) - _DECISION_FIELDS
        if unknown:
            raise ValueError(f"unknown decision field: {sorted(unknown)[0]}")
        if _contains_forbidden_model_evidence(record):
            raise ValueError("adjudication decision must not contain model-result paths or metrics")
        choice = record.get("decision")
        if choice not in _DECISIONS:
            raise ValueError("decision must be accept-a, accept-b, accept-both, replace, or exclude")
        if choice in {"accept-a", "accept-b", "accept-both"} and set(record) != {"decision"}:
            raise ValueError("accept decisions may contain only decision")
        if choice == "accept-a" and item.annotator_a_value is None:
            raise ValueError("accept-a requires an annotator A claim")
        if choice == "accept-b" and item.annotator_b_value is None:
            raise ValueError("accept-b requires an annotator B claim")
        if choice == "accept-both" and (
            item.annotator_a_value is None or item.annotator_b_value is None
        ):
            raise ValueError("accept-both requires claims from both annotators")
        if choice == "exclude" and set(record) != {"decision", "rationale"}:
            raise ValueError("exclude requires only decision and rationale")
        if choice == "replace":
            allowed = {"decision", "final_type", "final_value", "final_quote", "page_or_lines", "rationale"}
            if not {"decision", "final_type", "final_value", "final_quote", "rationale"} <= set(record) or set(record) - allowed:
                raise ValueError("replace requires final_type, final_value, final_quote, and rationale")
            final_type = _required_text(record, "final_type")
            allowed_types = (
                vocabulary["entity_types"]
                if item.claim_kind == "entity"
                else vocabulary["relationship_types"]
            )
            if final_type not in allowed_types:
                raise ValueError("replace final_type is not in the closed annotation vocabulary")
            base_value = item.annotator_a_value or item.annotator_b_value
            if (
                item.claim_kind == "entity"
                and final_type == "indicator"
                and (
                    not isinstance(base_value, Mapping)
                    or base_value.get("indicator_subtype")
                    not in vocabulary["indicator_subtypes"]
                )
            ):
                raise ValueError("indicator replacement requires a valid indicator subtype")
            final_value = _required_text(record, "final_value")
            if item.claim_kind == "relationship" and _RELATIONSHIP_ENDPOINTS.fullmatch(final_value) is None:
                raise ValueError("relationship final_value must be source_entity_id -> target_entity_id")
        if choice in {"replace", "exclude"}:
            _required_text(record, "rationale")


def _required_text(record: Mapping[str, object], field: str) -> str:
    value = record.get(field)
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field} is required")
    return unicodedata.normalize("NFC", value).strip()


def _replacement(
    item: Disagreement,
    decision: Mapping[str, object],
    a: EntityAnnotation | RelationshipAnnotation | None,
    b: EntityAnnotation | RelationshipAnnotation | None,
    comparison: Comparison,
) -> EntityAnnotation | RelationshipAnnotation:
    rationale = _required_text(decision, "rationale")
    del rationale
    base = a or b
    if base is None:
        raise ValueError("replace requires an existing claim")
    final_type = _required_text(decision, "final_type")
    final_value = _required_text(decision, "final_value")
    final_quote = _required_text(decision, "final_quote")
    page_or_lines = _required_text(decision, "page_or_lines") if decision.get("page_or_lines") not in (None, "") else base.page_or_lines
    if item.claim_kind == "entity":
        assert isinstance(base, EntityAnnotation)
        return replace(
            base,
            entity_type=final_type,
            indicator_subtype=(base.indicator_subtype if final_type == "indicator" else None),
            normalized_value=final_value,
            mention_as_written=final_value,
            supporting_quote=final_quote,
            page_or_lines=page_or_lines,
        )
    assert isinstance(base, RelationshipAnnotation)
    matched = _RELATIONSHIP_ENDPOINTS.fullmatch(final_value)
    assert matched is not None
    source_id, target_id = matched.group(1), matched.group(2)
    annotation = comparison.annotator_a if a is not None else comparison.annotator_b
    entity_ids = {
        entity.entity_id
        for entity in annotation.entities
        if entity.document_id == base.document_id
    }
    if source_id not in entity_ids or target_id not in entity_ids:
        raise ValueError("replacement relationship endpoint not found in document")
    return replace(
        base,
        source_entity_id=source_id,
        target_entity_id=target_id,
        relationship_type=final_type,
        supporting_quote=final_quote,
        page_or_lines=page_or_lines,
    )


def apply_adjudication(comparison: Comparison, decisions: Mapping[str, object] | list[Mapping[str, object]]) -> ReferenceSet:
    """Apply explicit human decisions; no annotation claim is silently corrected."""
    rows = _decision_rows(decisions)
    _validate_decision_rows(comparison, rows)
    a_map, b_map = _claim_maps(comparison)
    selected_entities: list[tuple[str, EntityAnnotation]] = [("a", item) for item in comparison.annotator_a.entities]
    selected_relationships: list[tuple[str, RelationshipAnnotation]] = [("a", item) for item in comparison.annotator_a.relationships]

    def remove(side: str, claim: EntityAnnotation | RelationshipAnnotation | None) -> None:
        if claim is None:
            return
        target = selected_entities if isinstance(claim, EntityAnnotation) else selected_relationships
        key = canonical_bytes(claim.to_dict())
        target[:] = [(item_side, item_claim) for item_side, item_claim in target if not (item_side == side and canonical_bytes(item_claim.to_dict()) == key)]

    def add(side: str, claim: EntityAnnotation | RelationshipAnnotation | None) -> None:
        if claim is None:
            return
        target = selected_entities if isinstance(claim, EntityAnnotation) else selected_relationships
        key = canonical_bytes(claim.to_dict())
        if not any(item_side == side and canonical_bytes(item_claim.to_dict()) == key for item_side, item_claim in target):
            target.append((side, claim))

    frozen_decisions: list[dict[str, object]] = []
    for item in comparison.disagreements:
        record = rows[item.disagreement_id]
        choice = record.get("decision")
        a = a_map.get(canonical_bytes(item.annotator_a_value)) if item.annotator_a_value else None
        b = b_map.get(canonical_bytes(item.annotator_b_value)) if item.annotator_b_value else None
        remove("a", a)
        remove("b", b)
        if choice == "accept-a":
            add("a", a)
        elif choice == "accept-b":
            add("b", b)
        elif choice == "accept-both":
            add("a", a)
            add("b", b)
        elif choice == "replace":
            add("a" if a is not None else "b", _replacement(item, record, a, b, comparison))
        frozen_decisions.append({"disagreement_id": item.disagreement_id, **dict(record)})

    return _reference_from_selected(comparison, tuple(sorted(frozen_decisions, key=lambda record: str(record["disagreement_id"]))), selected_entities, selected_relationships)


def _reference_from_selected(
    comparison: Comparison,
    decisions: tuple[dict[str, object], ...],
    selected_entities: list[tuple[str, EntityAnnotation]],
    selected_relationships: list[tuple[str, RelationshipAnnotation]],
) -> ReferenceSet:
    ordered_entities = sorted(selected_entities, key=lambda pair: (entity_key(pair[1]), pair[0], canonical_bytes(pair[1].to_dict())))
    remapped: list[EntityAnnotation] = []
    by_origin: dict[tuple[str, str, str], EntityAnnotation] = {}
    by_key: dict[tuple[str, str, str | None, str], EntityAnnotation] = {}
    for index, (side, entity) in enumerate(ordered_entities, start=1):
        result = replace(entity, entity_id=f"ref-e-{index:04d}")
        remapped.append(result)
        by_origin[(side, entity.document_id, entity.entity_id)] = result
        by_key.setdefault(entity_key(entity), result)

    source_sets = {"a": comparison.annotator_a, "b": comparison.annotator_b}
    remapped_relationships: list[RelationshipAnnotation] = []
    for index, (side, relationship) in enumerate(sorted(selected_relationships, key=lambda pair: (pair[0], canonical_bytes(pair[1].to_dict()))), start=1):
        source_entities = {(item.document_id, item.entity_id): item for item in source_sets.get(side, comparison.annotator_a).entities}
        source = by_origin.get((side, relationship.document_id, relationship.source_entity_id))
        target = by_origin.get((side, relationship.document_id, relationship.target_entity_id))
        if source is None and (relationship.document_id, relationship.source_entity_id) in source_entities:
            source = by_key.get(entity_key(source_entities[(relationship.document_id, relationship.source_entity_id)]))
        if target is None and (relationship.document_id, relationship.target_entity_id) in source_entities:
            target = by_key.get(entity_key(source_entities[(relationship.document_id, relationship.target_entity_id)]))
        if source is None or target is None:
            raise ValueError("adjudicated relationship endpoint is absent from reference")
        remapped_relationships.append(replace(relationship, relationship_id=f"ref-r-{index:04d}", source_entity_id=source.entity_id, target_entity_id=target.entity_id))
    return ReferenceSet(comparison, decisions, tuple(remapped), tuple(remapped_relationships))


def _utc_freeze_time(value: object) -> str:
    if not isinstance(value, str) or not value.endswith("Z"):
        raise ValueError("frozen_at must be a UTC ISO-8601 timestamp ending in Z")
    try:
        parsed = datetime.fromisoformat(value[:-1] + "+00:00")
    except ValueError as error:
        raise ValueError("frozen_at must be a UTC ISO-8601 timestamp ending in Z") from error
    if parsed.tzinfo is None or parsed.utcoffset() != timezone.utc.utcoffset(parsed):
        raise ValueError("frozen_at must be a UTC ISO-8601 timestamp ending in Z")
    return value


def freeze_reference(reference: ReferenceSet, output_root: str | Path, *, frozen_at: str | None = None) -> ReferenceFreeze:
    """Write exactly the six immutable, provenance-bound human reference files."""
    if frozen_at is None:
        frozen_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    frozen_at = _utc_freeze_time(frozen_at)
    documents = reference.comparison.annotator_a.document_ids
    if (
        len(documents) != 16
        or len(set(documents)) != 16
        or reference.comparison.annotator_b.document_ids != documents
    ):
        raise ValueError("reference freeze requires exact 16-document coverage")
    root = Path(output_root) / "annotations"
    if root.exists():
        raise ValueError("refusing to write into an existing annotations freeze directory")
    root.mkdir(parents=True)
    files = {
        "annotator-a.v1.json": reference.comparison.annotator_a.to_dict(),
        "annotator-b.v1.json": reference.comparison.annotator_b.to_dict(),
        "agreement.v1.json": reference.comparison.to_dict(),
        "adjudication.v1.json": {"schema_version": 1, "decisions": list(reference.decisions)},
        "reference.v1.json": reference.to_dict(),
    }
    hashes = {name: write_new_json(root / name, value) for name, value in files.items()}
    freeze = {
        "schema_version": 1,
        "annotator_a_workbook_sha256": reference.comparison.annotator_a.workbook_sha256,
        "annotator_b_workbook_sha256": reference.comparison.annotator_b.workbook_sha256,
        "imported_json_sha256": {name: hashes[name] for name in ("annotator-a.v1.json", "annotator-b.v1.json")},
        "adjudication_sha256": hashes["adjudication.v1.json"],
        "input_manifest_digest": reference.comparison.annotator_a.manifest_digest,
        "document_count": 16,
        "agreement_metrics": {
            "entities": reference.comparison.entity_metrics.to_dict(),
            "relationships": reference.comparison.relationship_metrics.to_dict(),
            "quotes": reference.comparison.quote_agreement.to_dict(),
        },
        "frozen_at": frozen_at,
    }
    hashes["reference-freeze.v1.json"] = write_new_json(root / "reference-freeze.v1.json", freeze)
    return ReferenceFreeze(str(root), 16, hashes["reference-freeze.v1.json"], frozen_at, hashes)


def _contains_forbidden_model_evidence(value: object) -> bool:
    if isinstance(value, Mapping):
        for key, child in value.items():
            lower = str(key).casefold()
            if any(token in lower for token in ("model", "result", "metric", "run", "condition", "tim")):
                return True
            if _contains_forbidden_model_evidence(child):
                return True
    elif isinstance(value, (list, tuple)):
        return any(_contains_forbidden_model_evidence(child) for child in value)
    elif isinstance(value, str):
        lower = value.casefold().replace("\\", "/")
        return "/runs/" in lower or "/results/" in lower or lower.startswith("runs/") or lower.startswith("results/")
    return False


_RELATIONSHIP_ENDPOINTS = re.compile(r"^([^\s]+)\s*->\s*([^\s]+)$")
