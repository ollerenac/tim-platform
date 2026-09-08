"""Build and check the blind two-document EXP-02 practice package.

Practice material is intentionally separate from the fixed final reference.  Its
two identifiers start with ``practice-`` and its documents have role ``practice``;
the normal final-workbook builder accepts neither shape.
"""

from __future__ import annotations

from dataclasses import dataclass
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path, PurePosixPath
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from .jsonio import _open_new_file, canonical_bytes
from .workbook import (
    DOCUMENT_COLUMNS,
    ENTITY_COLUMNS,
    MAX_ANNOTATION_ROWS,
    QUESTION_COLUMNS,
    RELATIONSHIP_COLUMNS,
    SHEET_NAMES,
    WorkbookSummary,
    _annotation_types,
    _assert_no_formulas,
    _load_workbook,
    _require_regular_xlsx,
    _validate_documents,
    _validate_doubts,
    _validate_entities,
    _validate_relationships,
    configure_annotation_sheet,
)


class PracticeError(ValueError):
    """The local practice package is incomplete or unsafe."""


@dataclass(frozen=True)
class PracticeDocument:
    """One historical document allowed only for annotator practice."""

    document_id: str
    practice_kind: str
    historical_stable_id: str
    exclusion_reason: str
    role: str
    source_id: str
    source_class: str
    title: str
    author: str
    published_at: str
    origin_url: str
    original_path: str
    text_path: str
    media_type: str
    word_count: int
    original_sha256: str
    text_sha256: str
    acquired_at: str


@dataclass(frozen=True)
class PracticeManifest:
    """Validated local provenance for the two practice documents."""

    digest: str
    documents: tuple[PracticeDocument, ...]


@dataclass(frozen=True)
class AnswerKey:
    """Completed rows kept out of the workbook delivered to annotators."""

    entities: tuple[dict[str, str | None], ...]
    relationships: tuple[dict[str, str | None], ...]


_PACKAGE_ROOT = Path(__file__).resolve().parents[2]
_SHA256 = re.compile(r"^[0-9a-f]{64}$")
_ID = re.compile(r"^practice-[a-z0-9-]+$")
_ENTITY_KEY_COLUMNS = [
    "document_id", "entity_id", "entity_type", "indicator_subtype",
    "normalized_value", "mention_as_written", "supporting_quote",
    "page_or_lines", "certainty", "notes",
]
_RELATIONSHIP_KEY_COLUMNS = [
    "document_id", "relationship_id", "source_entity_id", "relationship_type",
    "target_entity_id", "supporting_quote", "page_or_lines", "certainty", "notes",
]


def _read_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise PracticeError(f"cannot read JSON: {path.name}") from error
    if not isinstance(value, dict):
        raise PracticeError(f"JSON object required: {path.name}")
    return value


def _inside(root: Path, relative: object, field: str) -> Path:
    if not isinstance(relative, str) or not relative:
        raise PracticeError(f"{field} must be a non-empty local path")
    pure = PurePosixPath(relative)
    if pure.is_absolute() or ".." in pure.parts or pure.as_posix() != relative:
        raise PracticeError(f"{field} must be a safe relative path")
    target = root / pure
    if not target.is_file() or target.is_symlink():
        raise PracticeError(f"{field} must name a regular local file")
    return target


def _required_text(row: dict[str, Any], field: str) -> str:
    value = row.get(field)
    if not isinstance(value, str) or not value.strip():
        raise PracticeError(f"practice document {field} is required")
    return value


def _document(row: object, root: Path) -> PracticeDocument:
    if not isinstance(row, dict):
        raise PracticeError("practice document must be an object")
    expected = {field.name for field in PracticeDocument.__dataclass_fields__.values()}
    if set(row) != expected:
        raise PracticeError("practice document fields do not match schema")
    document = PracticeDocument(**row)
    if not _ID.fullmatch(document.document_id):
        raise PracticeError("practice document_id must begin with practice-")
    if document.practice_kind not in {"narrative", "indicator-heavy"}:
        raise PracticeError("practice_kind must be narrative or indicator-heavy")
    if document.role != "practice":
        raise PracticeError("practice document role must be practice")
    if document.exclusion_reason not in {"previously_evaluated", "development_material"}:
        raise PracticeError("practice exclusion reason is not allowed")
    if document.source_class not in {"institutional", "technical-research"}:
        raise PracticeError("practice source class is invalid")
    if document.media_type not in {"html", "pdf"}:
        raise PracticeError("practice media type is invalid")
    if not isinstance(document.word_count, int) or not 500 <= document.word_count <= 50_000:
        raise PracticeError("practice word_count is outside the allowed range")
    for field in (
        "historical_stable_id", "source_id", "title", "author", "published_at",
        "origin_url", "acquired_at",
    ):
        _required_text(row, field)
    if not document.origin_url.startswith("https://"):
        raise PracticeError("practice origin_url must be HTTPS")
    for field, value in (("original_sha256", document.original_sha256), ("text_sha256", document.text_sha256)):
        if not isinstance(value, str) or not _SHA256.fullmatch(value):
            raise PracticeError(f"{field} must be a lowercase SHA-256")
    original = _inside(root, document.original_path, "original_path")
    converted = _inside(root, document.text_path, "text_path")
    if hashlib.sha256(original.read_bytes()).hexdigest() != document.original_sha256:
        raise PracticeError("practice original hash does not match")
    if hashlib.sha256(converted.read_bytes()).hexdigest() != document.text_sha256:
        raise PracticeError("practice converted-text hash does not match")
    return document


def load_practice_manifest(practice_root: str | Path) -> PracticeManifest:
    """Load exactly two local, excluded documents and bind their copied bytes."""
    root = Path(practice_root)
    raw = _read_json(root / "manifest.v1.json")
    if set(raw) != {"schema_version", "documents"} or raw.get("schema_version") != 1:
        raise PracticeError("practice manifest schema is invalid")
    rows = raw.get("documents")
    if not isinstance(rows, list) or len(rows) != 2:
        raise PracticeError("practice manifest must contain exactly two documents")
    documents = tuple(_document(row, root) for row in rows)
    if len({document.document_id for document in documents}) != 2:
        raise PracticeError("practice manifest has duplicate document IDs")
    if {document.practice_kind for document in documents} != {"narrative", "indicator-heavy"}:
        raise PracticeError("practice manifest needs one narrative and one indicator-heavy document")

    exclusions = _read_json(_PACKAGE_ROOT / "config" / "exclusions.v1.json")
    exclusion_rows = exclusions.get("entries")
    if not isinstance(exclusion_rows, list):
        raise PracticeError("exclusion ledger is invalid")
    excluded = {
        entry.get("stable_id"): entry
        for entry in exclusion_rows
        if isinstance(entry, dict) and isinstance(entry.get("stable_id"), str)
    }
    for document in documents:
        historical = excluded.get(document.historical_stable_id)
        if not isinstance(historical, dict):
            raise PracticeError("practice historical input is not in the exclusion ledger")
        if historical.get("reason") != document.exclusion_reason:
            raise PracticeError("practice exclusion reason does not match ledger")
        if historical.get("sha256") != document.original_sha256:
            raise PracticeError("practice original hash does not match exclusion ledger")
    return PracticeManifest(
        digest=hashlib.sha256(canonical_bytes(raw)).hexdigest(), documents=documents,
    )


def _validation(sheet, column: str, formula: str) -> None:  # type: ignore[no-untyped-def]
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.showErrorMessage = True
    validation.error = "Seleccione un valor de la lista cerrada de EXP-02."
    sheet.add_data_validation(validation)
    validation.add(f"{column}2:{column}{MAX_ANNOTATION_ROWS + 1}")


def _safe_save_new(workbook: Workbook, output: Path) -> None:
    """Serialize first, then create the target once through a no-follow descriptor."""
    buffer = BytesIO()
    workbook.save(buffer)
    try:
        descriptor = _open_new_file(output)
        with os.fdopen(descriptor, "xb") as handle:
            handle.write(buffer.getvalue())
    except FileExistsError as error:
        raise PracticeError("refusing to overwrite practice workbook") from error
    except ValueError as error:
        raise PracticeError(str(error)) from error
    except OSError as error:
        raise PracticeError(f"cannot write practice workbook: {output}") from error


def build_practice_workbook(practice_root: str | Path, output_path: str | Path) -> str:
    """Create a blank, LibreOffice-compatible exercise workbook without answer rows."""
    manifest = load_practice_manifest(practice_root)
    output = Path(output_path)
    if output.name != "practice.xlsx":
        raise PracticeError("practice output filename must be practice.xlsx")
    if output.is_symlink():
        raise PracticeError("practice output must not be a symlink")

    vocabulary = _read_json(_PACKAGE_ROOT / "config" / "annotation-types.v1.json")
    workbook = Workbook()
    readme = workbook.active
    readme.title = "LEEME"
    documents = workbook.create_sheet("DOCUMENTOS")
    entities = workbook.create_sheet("ENTIDADES")
    relationships = workbook.create_sheet("RELACIONES")
    doubts = workbook.create_sheet("DUDAS")
    lists = workbook.create_sheet("_LISTAS")

    readme["A1"] = "EXP-02 practice workbook"
    readme["A2"] = "role"
    readme["B2"] = "practice"
    readme["A3"] = "practice_manifest_sha256"
    readme["B3"] = manifest.digest
    readme["A4"] = "document_scope"
    readme["B4"] = "two-excluded-documents"
    readme["A5"] = "Ejercicio ciego: complete las filas sin consultar la clave de respuestas."
    readme["A6"] = "Estos dos documentos son solo práctica; no pueden entrar en la referencia final."
    readme.protection.sheet = True
    for row in readme.iter_rows():
        for cell in row:
            cell.protection = Protection(locked=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 38
    readme.column_dimensions["B"].width = 82

    configure_annotation_sheet(documents, DOCUMENT_COLUMNS, unlock=False)
    for row, document in enumerate(manifest.documents, start=2):
        values = [
            document.document_id, document.source_id, document.title, document.author,
            document.published_at, document.origin_url, document.text_path, document.media_type,
            document.word_count, None, None,
        ]
        for column, value in enumerate(values, start=1):
            cell = documents.cell(row, column, value)
            cell.alignment = Alignment(wrap_text=column in {3, 6, 7, 11}, vertical="top")
            cell.protection = Protection(locked=column not in {10, 11})

    for sheet, columns in ((entities, ENTITY_COLUMNS), (relationships, RELATIONSHIP_COLUMNS), (doubts, QUESTION_COLUMNS)):
        configure_annotation_sheet(sheet, columns, unlock=True)

    list_values = [
        ("entity_types", vocabulary["entity_types"]),
        ("relationship_types", vocabulary["relationship_types"]),
        ("certainty_values", vocabulary["certainty_values"]),
        ("indicator_subtypes", vocabulary["indicator_subtypes"]),
        ("practice_document_ids", [document.document_id for document in manifest.documents]),
    ]
    for column, (name, values) in enumerate(list_values, start=1):
        lists.cell(1, column, name)
        for row, value in enumerate(values, start=2):
            lists.cell(row, column, value)
    lists.sheet_state = "hidden"
    lists.protection.sheet = True
    _validation(entities, "A", "'_LISTAS'!$E$2:$E$3")
    _validation(entities, "C", "'_LISTAS'!$A$2:$A$12")
    _validation(entities, "D", "'_LISTAS'!$D$2:$D$8")
    _validation(entities, "I", "'_LISTAS'!$C$2:$C$3")
    _validation(relationships, "A", "'_LISTAS'!$E$2:$E$3")
    _validation(relationships, "D", "'_LISTAS'!$B$2:$B$6")
    _validation(relationships, "H", "'_LISTAS'!$C$2:$C$3")
    _validation(doubts, "A", "'_LISTAS'!$E$2:$E$3")
    _safe_save_new(workbook, output)
    return str(output)


def validate_practice_workbook(
    path: str | Path, practice_root: str | Path,
) -> WorkbookSummary:
    """Validate either the blank exercise or an edited practice workbook."""
    manifest = load_practice_manifest(practice_root)
    output = _require_regular_xlsx(path)
    if output.name != "practice.xlsx":
        raise PracticeError("practice workbook filename must be practice.xlsx")
    workbook = _load_workbook(output)
    if workbook.sheetnames != SHEET_NAMES:
        raise PracticeError("practice workbook sheets do not match the fixed layout")
    if workbook["_LISTAS"].sheet_state != "hidden":
        raise PracticeError("practice vocabulary sheet must be hidden")
    readme = workbook["LEEME"]
    if (
        readme["B2"].value != "practice"
        or readme["B3"].value != manifest.digest
        or readme["B4"].value != "two-excluded-documents"
    ):
        raise PracticeError("practice workbook identity is invalid")
    if (
        not readme.protection.sheet
        or not readme["B2"].protection.locked
        or not readme["B3"].protection.locked
        or not readme["B4"].protection.locked
    ):
        raise PracticeError("practice identity cells must be protected")
    _assert_no_formulas(workbook)
    documents = manifest.documents
    _validate_documents(workbook["DOCUMENTOS"], documents)  # type: ignore[arg-type]
    document_ids = {document.document_id for document in documents}
    vocabulary = _annotation_types()
    entities = _validate_entities(workbook["ENTIDADES"], document_ids, vocabulary)
    relationship_count = _validate_relationships(
        workbook["RELACIONES"], document_ids, entities, vocabulary,
    )
    question_count = _validate_doubts(workbook["DUDAS"], document_ids)
    return WorkbookSummary(
        manifest_digest=manifest.digest,
        annotator_id="practice",
        document_count=2,
        entity_count=len(entities),
        relationship_count=relationship_count,
        question_count=question_count,
    )


def _table_rows(markdown: str, heading: str, columns: list[str]) -> tuple[dict[str, str | None], ...]:
    match = re.search(rf"(?ms)^## {re.escape(heading)}\s*$\n(.*?)(?=^## |\Z)", markdown)
    if match is None:
        raise PracticeError(f"answer key is missing {heading} table")
    lines = [line.strip() for line in match.group(1).splitlines() if line.strip().startswith("|")]
    if len(lines) < 3:
        raise PracticeError(f"answer key {heading} table is empty")
    parsed = [[cell.strip() for cell in line.strip("|").split("|")] for line in lines]
    if parsed[0] != columns or any(not re.fullmatch(r"[-: ]+", cell) for cell in parsed[1]):
        raise PracticeError(f"answer key {heading} columns are invalid")
    result: list[dict[str, str | None]] = []
    for row in parsed[2:]:
        if len(row) != len(columns):
            raise PracticeError(f"answer key {heading} row has wrong width")
        result.append({name: (None if value in {"", "-"} else value) for name, value in zip(columns, row)})
    return tuple(result)


def read_answer_key(path: str | Path) -> AnswerKey:
    """Parse and validate answer rows without ever copying them into practice.xlsx."""
    try:
        markdown = Path(path).read_text(encoding="utf-8")
    except OSError as error:
        raise PracticeError("answer key is unavailable") from error
    entities = _table_rows(markdown, "Entidades", _ENTITY_KEY_COLUMNS)
    relationships = _table_rows(markdown, "Relaciones", _RELATIONSHIP_KEY_COLUMNS)
    manifest = load_practice_manifest(_PACKAGE_ROOT / "practice")
    vocabulary = _read_json(_PACKAGE_ROOT / "config" / "annotation-types.v1.json")
    document_ids = {document.document_id for document in manifest.documents}
    local_entities: set[tuple[str, str]] = set()
    for entity in entities:
        document_id = entity["document_id"]
        entity_id = entity["entity_id"]
        if document_id not in document_ids or not entity_id:
            raise PracticeError("answer-key entity has an unknown document or empty identifier")
        if (document_id, entity_id) in local_entities:
            raise PracticeError("answer-key entity identifier is duplicated")
        local_entities.add((document_id, entity_id))
        entity_type = entity["entity_type"]
        if entity_type not in vocabulary["entity_types"]:
            raise PracticeError("answer-key entity type is outside closed vocabulary")
        subtype = entity["indicator_subtype"]
        if entity_type == "indicator" and subtype not in vocabulary["indicator_subtypes"]:
            raise PracticeError("answer-key indicator subtype is outside closed vocabulary")
        if entity_type != "indicator" and subtype is not None:
            raise PracticeError("answer-key subtype belongs only to an indicator")
        if any(entity[field] is None for field in ("normalized_value", "mention_as_written", "supporting_quote", "page_or_lines", "certainty")):
            raise PracticeError("answer-key entity has a required empty value")
    for relationship in relationships:
        document_id = relationship["document_id"]
        if document_id not in document_ids or relationship["relationship_type"] not in vocabulary["relationship_types"]:
            raise PracticeError("answer-key relationship is outside closed vocabulary")
        if (document_id, relationship["source_entity_id"] or "") not in local_entities or (document_id, relationship["target_entity_id"] or "") not in local_entities:
            raise PracticeError("answer-key relationship endpoint is missing in its document")
        if any(relationship[field] is None for field in ("relationship_id", "supporting_quote", "page_or_lines", "certainty")):
            raise PracticeError("answer-key relationship has a required empty value")
    return AnswerKey(entities=entities, relationships=relationships)
