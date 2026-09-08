"""Import immutable human annotation records without changing their claims."""

from __future__ import annotations

from collections.abc import Mapping
from dataclasses import asdict, dataclass
import hashlib
import os
from pathlib import Path
import re
import stat
import tempfile
import unicodedata

from openpyxl import load_workbook

from .jsonio import canonical_bytes, sha256_file
from .workbook import ENTITY_COLUMNS, RELATIONSHIP_COLUMNS, WorkbookError, validate_workbook


class AnnotationError(ValueError):
    """A validated annotation workbook cannot be imported faithfully."""


_PAGE_OR_LINES = re.compile(r"^(?:page|pages|line|lines)\s+\d+(?:\s*-\s*\d+)?$", re.IGNORECASE)


@dataclass(frozen=True)
class EntityAnnotation:
    document_id: str
    entity_id: str
    entity_type: str
    indicator_subtype: str | None
    normalized_value: str
    mention_as_written: str
    supporting_quote: str
    page_or_lines: str
    certainty: str
    notes: str | None

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


@dataclass(frozen=True)
class RelationshipAnnotation:
    document_id: str
    relationship_id: str
    source_entity_id: str
    relationship_type: str
    target_entity_id: str
    supporting_quote: str
    page_or_lines: str
    certainty: str
    notes: str | None

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


@dataclass(frozen=True)
class AnnotationSet:
    annotator_id: str
    manifest_digest: str
    workbook_sha256: str
    entities: tuple[EntityAnnotation, ...]
    relationships: tuple[RelationshipAnnotation, ...]
    document_ids: tuple[str, ...] = ()
    document_minutes: tuple[tuple[str, int], ...] = ()
    workbook_path: str | None = None

    def to_dict(self) -> dict[str, object]:
        return {
            "schema_version": 1,
            "annotator_id": self.annotator_id,
            "input_manifest_digest": self.manifest_digest,
            "workbook_sha256": self.workbook_sha256,
            "entities": [entity.to_dict() for entity in self.entities],
            "relationships": [relationship.to_dict() for relationship in self.relationships],
            "document_ids": list(self.document_ids),
            "document_minutes": {
                document_id: minutes for document_id, minutes in self.document_minutes
            },
            "workbook_path": self.workbook_path,
        }

    @property
    def digest(self) -> str:
        return hashlib.sha256(canonical_bytes(self.to_dict())).hexdigest()


def _text(value: object, *, required: bool, field: str) -> str | None:
    if value is None:
        if required:
            raise AnnotationError(f"{field} is required")
        return None
    if not isinstance(value, str):
        raise AnnotationError(f"{field} must be text")
    normalized = unicodedata.normalize("NFC", value).strip()
    if required and not normalized:
        raise AnnotationError(f"{field} is required")
    return normalized or None


def _row_values(sheet, row: int, columns: list[str]) -> dict[str, object]:
    return {column: sheet.cell(row, index).value for index, column in enumerate(columns, start=1)}


def _has_value(values: Mapping[str, object]) -> bool:
    return any(value not in (None, "") for value in values.values())


def _page_or_lines(value: object, field: str) -> str:
    result = _text(value, required=True, field=field)
    assert result is not None
    if not _PAGE_OR_LINES.fullmatch(result):
        raise AnnotationError(f"{field} must be a page or line reference")
    return result


def _snapshot_workbook(path: str | Path) -> Path:
    """Copy the opened regular workbook once so validation and import see one file."""
    flags = os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0)
    try:
        descriptor = os.open(path, flags)
    except OSError as error:
        raise AnnotationError("annotation workbook is unavailable") from error
    snapshot_dir = Path(tempfile.mkdtemp(prefix="exp02-annotation-"))
    snapshot = snapshot_dir / Path(path).name
    try:
        if not stat.S_ISREG(os.fstat(descriptor).st_mode):
            raise AnnotationError("annotation workbook must be a regular file")
        with os.fdopen(descriptor, "rb") as source, snapshot.open("xb") as destination:
            descriptor = -1
            while block := source.read(1024 * 1024):
                destination.write(block)
            return snapshot
    except BaseException:
        snapshot.unlink(missing_ok=True)
        snapshot_dir.rmdir()
        raise
    finally:
        if descriptor != -1:
            os.close(descriptor)


def import_annotations(path: str | Path, manifest: Mapping[str, object] | str | Path) -> AnnotationSet:
    """Validate then import a human workbook, preserving every stated claim.

    Workbook validation protects its closed vocabulary and document foreign keys;
    this layer makes normalized, immutable records without applying semantic
    correction to the annotator's values.
    """
    snapshot = _snapshot_workbook(path)
    try:
        try:
            summary = validate_workbook(snapshot, manifest)
        except WorkbookError as error:
            raise AnnotationError(str(error)) from error
        if summary.document_count != 16:
            raise AnnotationError("annotation workbook must cover exactly 16 documents")
        try:
            workbook = load_workbook(snapshot, data_only=False, keep_links=False)
        except (OSError, ValueError) as error:
            raise AnnotationError("annotation workbook cannot be read") from error

        entities: list[EntityAnnotation] = []
        for row in range(2, workbook["ENTIDADES"].max_row + 1):
            values = _row_values(workbook["ENTIDADES"], row, ENTITY_COLUMNS)
            if not _has_value(values):
                continue
            entities.append(EntityAnnotation(
                document_id=_text(values["document_id"], required=True, field="document_id") or "",
                entity_id=_text(values["entity_id"], required=True, field="entity_id") or "",
                entity_type=_text(values["entity_type"], required=True, field="entity_type") or "",
                indicator_subtype=_text(values["indicator_subtype"], required=False, field="indicator_subtype"),
                normalized_value=_text(values["normalized_value"], required=True, field="normalized_value") or "",
                mention_as_written=_text(values["mention_as_written"], required=True, field="mention_as_written") or "",
                supporting_quote=_text(values["supporting_quote"], required=True, field="supporting_quote") or "",
                page_or_lines=_page_or_lines(values["page_or_lines"], "page_or_lines"),
                certainty=_text(values["certainty"], required=True, field="certainty") or "",
                notes=_text(values["notes"], required=False, field="notes"),
            ))
        entities_by_id = {(entity.document_id, entity.entity_id): entity for entity in entities}

        relationships: list[RelationshipAnnotation] = []
        for row in range(2, workbook["RELACIONES"].max_row + 1):
            values = _row_values(workbook["RELACIONES"], row, RELATIONSHIP_COLUMNS)
            if not _has_value(values):
                continue
            document_id = _text(values["document_id"], required=True, field="document_id") or ""
            source = _text(values["source_entity_id"], required=True, field="source_entity_id") or ""
            target = _text(values["target_entity_id"], required=True, field="target_entity_id") or ""
            if (document_id, source) not in entities_by_id or (document_id, target) not in entities_by_id:
                raise AnnotationError("endpoint not found in document")
            relationships.append(RelationshipAnnotation(
                document_id=document_id,
                relationship_id=_text(values["relationship_id"], required=True, field="relationship_id") or "",
                source_entity_id=source,
                relationship_type=_text(values["relationship_type"], required=True, field="relationship_type") or "",
                target_entity_id=target,
                supporting_quote=_text(values["supporting_quote"], required=True, field="supporting_quote") or "",
                page_or_lines=_page_or_lines(values["page_or_lines"], "page_or_lines"),
                certainty=_text(values["certainty"], required=True, field="certainty") or "",
                notes=_text(values["notes"], required=False, field="notes"),
            ))
        document_ids = tuple(
            str(workbook["DOCUMENTOS"].cell(row, 1).value)
            for row in range(2, workbook["DOCUMENTOS"].max_row + 1)
            if workbook["DOCUMENTOS"].cell(row, 1).value not in (None, "")
        )
        minutes: list[tuple[str, int]] = []
        for row in range(2, workbook["DOCUMENTOS"].max_row + 1):
            document_id = workbook["DOCUMENTOS"].cell(row, 1).value
            if document_id in (None, ""):
                continue
            active_minutes = workbook["DOCUMENTOS"].cell(row, 10).value
            if (
                isinstance(active_minutes, bool)
                or not isinstance(active_minutes, int)
                or active_minutes <= 0
            ):
                raise AnnotationError(
                    "active minutes must be a positive integer for every document"
                )
            minutes.append((str(document_id), active_minutes))
        document_minutes = tuple(minutes)
        return AnnotationSet(
            annotator_id=summary.annotator_id,
            manifest_digest=summary.manifest_digest,
            workbook_sha256=sha256_file(snapshot),
            entities=tuple(entities),
            relationships=tuple(relationships),
            document_ids=document_ids,
            document_minutes=document_minutes,
        )
    finally:
        snapshot.unlink(missing_ok=True)
        snapshot.parent.rmdir()
