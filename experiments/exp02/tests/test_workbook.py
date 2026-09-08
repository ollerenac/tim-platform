"""Behavioral contracts for independent EXP-02 annotation workbooks."""

from __future__ import annotations

import hashlib
from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from exp02.jsonio import canonical_bytes
from exp02.workbook import WorkbookError, build_workbook, validate_workbook


ENTITY_COLUMNS = [
    "document_id",
    "entity_id",
    "entity_type",
    "indicator_subtype",
    "normalized_value",
    "mention_as_written",
    "supporting_quote",
    "page_or_lines",
    "certainty",
    "notes",
]
RELATIONSHIP_COLUMNS = [
    "document_id",
    "relationship_id",
    "source_entity_id",
    "relationship_type",
    "target_entity_id",
    "supporting_quote",
    "page_or_lines",
    "certainty",
    "notes",
]


def sample_manifest() -> dict[str, object]:
    def document(number: int) -> dict[str, object]:
        document_id = "doc-primary" if number == 1 else f"doc-{number:02d}"
        return {
        "document_id": document_id,
        "source_id": "source-one",
        "source_class": "institutional",
        "role": "primary",
        "title": f"Primary report {number}",
        "author": "Analyst",
        "author_basis": "feed_author",
        "published_at": "2026-08-30T10:00:00Z",
        "origin_url": f"https://example.test/{document_id}",
        "original_path": f"experiments/exp02/evidence/inputs/{document_id}/original.html",
        "text_path": f"experiments/exp02/evidence/inputs/{document_id}/converted.txt",
        "media_type": "html",
        "word_count": 600,
        "original_sha256": "a" * 64,
        "text_sha256": "b" * 64,
        "acquired_at": "2026-08-30T11:00:00Z",
        }
    return {
        "schema_version": 1,
        "policy_digest": "c" * 64,
        "candidate_ledger_digest": "d" * 64,
        "exclusion_ledger_digest": "e" * 64,
        "sources": [
            {
                "source_id": "source-one",
                "source_class": "institutional",
                "feed_url": "https://example.test/feed",
            }
        ],
        "documents": [document(number) for number in range(1, 17)],
    }


def sample_digest() -> str:
    return hashlib.sha256(canonical_bytes(sample_manifest())).hexdigest()


def _workbook(tmp_path: Path) -> Path:
    path = tmp_path / "annotator-a.xlsx"
    build_workbook(sample_manifest(), "annotator-a", path)
    return path


def _save_changed(path: Path, change) -> None:  # type: ignore[no-untyped-def]
    workbook = load_workbook(path)
    change(workbook)
    workbook.save(path)


def test_workbook_has_fixed_sheets_hidden_vocab_and_annotation_layout(tmp_path: Path) -> None:
    path = _workbook(tmp_path)

    workbook = load_workbook(path)

    assert workbook.sheetnames == [
        "LEEME", "DOCUMENTOS", "ENTIDADES", "RELACIONES", "DUDAS", "_LISTAS"
    ]
    assert workbook["_LISTAS"].sheet_state == "hidden"
    assert workbook["LEEME"]["B2"].value == "annotator-a"
    assert [cell.value for cell in workbook["ENTIDADES"][1]] == ENTITY_COLUMNS
    assert [cell.value for cell in workbook["RELACIONES"][1]] == RELATIONSHIP_COLUMNS
    assert [row[0].value for row in workbook["DOCUMENTOS"].iter_rows(min_row=2)] == [
        "doc-primary",
        *[f"doc-{number:02d}" for number in range(2, 17)],
    ]
    assert workbook["ENTIDADES"].freeze_panes == "A2"
    assert workbook["ENTIDADES"].auto_filter.ref == "A1:J1001"
    assert workbook["ENTIDADES"]["G2"].alignment.wrap_text is True
    assert workbook["ENTIDADES"]["J2"].alignment.wrap_text is True
    assert workbook["DOCUMENTOS"].sheet_format.defaultRowHeight >= 30
    assert workbook["ENTIDADES"].sheet_format.defaultRowHeight >= 30
    assert workbook["RELACIONES"].sheet_format.defaultRowHeight >= 30
    assert workbook["DUDAS"].sheet_format.defaultRowHeight >= 30
    assert workbook["LEEME"].protection.sheet is True
    assert workbook["LEEME"]["B2"].protection.locked is True
    assert workbook["ENTIDADES"]["A2"].protection.locked is False


def test_workbook_binds_manifest_digest_and_validates_clean_file(tmp_path: Path) -> None:
    path = _workbook(tmp_path)

    summary = validate_workbook(path, sample_manifest())

    assert summary.manifest_digest == sample_digest()
    assert summary.annotator_id == "annotator-a"
    assert summary.document_count == 16
    assert summary.entity_count == 0
    assert summary.relationship_count == 0


def test_workbook_rejects_legacy_extension_scope_and_noncanonical_filename(
    tmp_path: Path,
) -> None:
    manifest = sample_manifest()
    manifest["documents"][0]["role"] = "extension"  # type: ignore[index]

    with pytest.raises(WorkbookError, match="primary"):
        build_workbook(manifest, "annotator-a", tmp_path / "annotator-a.xlsx")
    with pytest.raises(WorkbookError, match="filename"):
        build_workbook(
            sample_manifest(), "annotator-a", tmp_path / "annotator-a-24.xlsx"
        )


def test_workbook_write_is_exclusive_and_rejects_symlink_parent(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "exclusive")
    original = path.read_bytes()
    with pytest.raises(WorkbookError, match="overwrite"):
        build_workbook(sample_manifest(), "annotator-a", path)
    assert path.read_bytes() == original

    outside = tmp_path / "outside"
    outside.mkdir()
    linked = tmp_path / "linked"
    linked.symlink_to(outside, target_is_directory=True)
    with pytest.raises(WorkbookError, match="symlink"):
        build_workbook(sample_manifest(), "annotator-a", linked / "annotator-a.xlsx")
    assert list(outside.iterdir()) == []


def test_workbook_rejects_invalid_annotation_edits(tmp_path: Path) -> None:
    cases = [
        (
            "duplicate entity ID",
            lambda workbook: (
                workbook["ENTIDADES"].append(
                    [
                        "doc-primary", "entity-1", "indicator", "ip", "192.0.2.1",
                        "192.0.2.1", "Observed 192.0.2.1.", "lines 1-1", "clear", None,
                    ]
                ),
                workbook["ENTIDADES"].append(
                    [
                        "doc-primary", "entity-1", "indicator", "ip", "198.51.100.1",
                        "198.51.100.1", "Observed 198.51.100.1.", "lines 2-2", "clear", None,
                    ]
                ),
            ),
            "duplicate entity_id",
        ),
        (
            "unknown document ID",
            lambda workbook: workbook["ENTIDADES"].append(
                [
                    "unknown-doc", "entity-1", "indicator", "ip", "192.0.2.1", "192.0.2.1",
                    "Observed 192.0.2.1.", "lines 1-1", "clear", None,
                ]
            ),
            "unknown document_id",
        ),
        (
            "missing quote",
            lambda workbook: workbook["ENTIDADES"].append(
                ["doc-primary", "entity-1", "indicator", "ip", "192.0.2.1", "192.0.2.1", None,
                 "lines 1-1", "clear", None]
            ),
            "supporting_quote",
        ),
        (
            "invalid closed type",
            lambda workbook: workbook["ENTIDADES"].append(
                [
                    "doc-primary", "entity-1", "made-up", None, "Thing", "Thing", "Thing appeared.",
                    "lines 1-1", "clear", None,
                ]
            ),
            "invalid entity_type",
        ),
        (
            "missing relation endpoint",
            lambda workbook: workbook["RELACIONES"].append(
                [
                    "doc-primary", "relationship-1", "missing-source", "uses", "missing-target",
                    "The report says it uses it.", "lines 1-1", "clear", None,
                ]
            ),
            "endpoint not found",
        ),
        (
            "formula in user cell",
            lambda workbook: setattr(workbook["ENTIDADES"]["E2"], "value", "=1+1"),
            "formula",
        ),
        (
            "changed digest",
            lambda workbook: setattr(workbook["LEEME"]["B3"], "value", "0" * 64),
            "manifest digest",
        ),
        (
            "wrong annotator identity",
            lambda workbook: setattr(workbook["LEEME"]["B2"], "value", "annotator-b"),
            "annotator identity",
        ),
    ]

    for index, (_name, change, expected) in enumerate(cases):
        path = _workbook(tmp_path / str(index))
        _save_changed(path, change)
        with pytest.raises(WorkbookError, match=expected):
            validate_workbook(path, sample_manifest())


def test_workbook_rejects_macros_and_external_links(tmp_path: Path) -> None:
    macro_path = _workbook(tmp_path / "macro")
    with ZipFile(macro_path, "a", ZIP_DEFLATED) as archive:
        archive.writestr("xl/vbaProject.bin", b"not a macro")
    with pytest.raises(WorkbookError, match="macros"):
        validate_workbook(macro_path, sample_manifest())

    link_path = _workbook(tmp_path / "link")
    with ZipFile(link_path, "a", ZIP_DEFLATED) as archive:
        archive.writestr("xl/externalLinks/externalLink1.xml", b"<externalLink/>")
    with pytest.raises(WorkbookError, match="external links"):
        validate_workbook(link_path, sample_manifest())

    hyperlink_path = _workbook(tmp_path / "cell-hyperlink")
    _save_changed(
        hyperlink_path,
        lambda workbook: setattr(
            workbook["LEEME"]["A5"], "hyperlink", "https://external.example.test/evidence"
        ),
    )
    with ZipFile(hyperlink_path) as archive:
        relationships = archive.read("xl/worksheets/_rels/sheet1.xml.rels")
    assert b'TargetMode="External"' in relationships
    with pytest.raises(WorkbookError, match="external links"):
        validate_workbook(hyperlink_path, sample_manifest())


def test_workbook_rejects_manifest_mismatch(tmp_path: Path) -> None:
    path = _workbook(tmp_path)
    altered_manifest = deepcopy(sample_manifest())
    altered_manifest["documents"][0]["title"] = "Altered after freeze"  # type: ignore[index]

    with pytest.raises(WorkbookError, match="manifest digest"):
        validate_workbook(path, altered_manifest)


def test_workbook_requires_exactly_sixteen_manifest_documents(tmp_path: Path) -> None:
    manifest = sample_manifest()
    manifest["documents"] = manifest["documents"][:-1]  # type: ignore[index]

    with pytest.raises(WorkbookError, match="exactly 16"):
        build_workbook(manifest, "annotator-a", tmp_path / "annotator-a.xlsx")
