"""Contracts for importing and freezing independent human annotations."""

from __future__ import annotations

from pathlib import Path
import os

import pytest
from openpyxl import load_workbook

import exp02.annotations as annotations
from exp02.annotations import AnnotationError, import_annotations
from exp02.jsonio import sha256_file
from exp02.workbook import build_workbook


def manifest() -> dict[str, object]:
    def document(document_id: str) -> dict[str, object]:
        return {
            "document_id": document_id,
            "source_id": "source-one",
            "source_class": "institutional",
            "role": "primary",
            "title": f"Report {document_id}",
            "author": "Analyst",
            "author_basis": "feed_author",
            "published_at": "2026-08-30T10:00:00Z",
            "origin_url": f"https://example.test/{document_id}",
            "original_path": f"experiments/exp02/evidence/inputs/{document_id}/original.html",
            "text_path": f"experiments/exp02/evidence/inputs/{document_id}/converted.txt",
            "media_type": "html",
            "word_count": 600,
            "original_sha256": "a" * 64,
            "text_sha256": "b" * 64,
            "acquired_at": "2026-08-30T11:00:00Z",
        }

    return {
        "schema_version": 1,
        "policy_digest": "c" * 64,
        "candidate_ledger_digest": "d" * 64,
        "exclusion_ledger_digest": "e" * 64,
        "sources": [{"source_id": "source-one", "source_class": "institutional", "feed_url": "https://example.test/feed"}],
        "documents": [
            document("doc-one"),
            document("doc-two"),
            *[document(f"doc-{number:02d}") for number in range(3, 17)],
        ],
    }


def workbook_fixture(tmp_path: Path) -> Path:
    path = tmp_path / "annotator-a.xlsx"
    build_workbook(manifest(), "annotator-a", path)
    workbook = load_workbook(path)
    for column, value in enumerate([
        "doc-one", "e001", "malware", None, " ExampleLoader ",
        " ExampleLoader ", " ExampleLoader contacted 192.0.2.4. ", " lines 1-1 ", "clear", None,
    ], start=1):
        workbook["ENTIDADES"].cell(2, column, value)
    for row in range(2, workbook["DOCUMENTOS"].max_row + 1):
        workbook["DOCUMENTOS"].cell(row, 10, 30)
    workbook.save(path)
    return path


def workbook_with_cross_doc_relation(tmp_path: Path) -> Path:
    path = workbook_fixture(tmp_path)
    workbook = load_workbook(path)
    for column, value in enumerate([
        "doc-two", "e002", "indicator", "ip", "192.0.2.4", "192.0.2.4",
        "Observed 192.0.2.4.", "lines 1-1", "clear", None,
    ], start=1):
        workbook["ENTIDADES"].cell(3, column, value)
    for column, value in enumerate([
        "doc-one", "r001", "e001", "uses", "e002", "ExampleLoader contacted 192.0.2.4.",
        "lines 1-1", "clear", None,
    ], start=1):
        workbook["RELACIONES"].cell(2, column, value)
    workbook.save(path)
    return path


def test_import_preserves_quotes_and_local_ids(tmp_path: Path) -> None:
    result = import_annotations(workbook_fixture(tmp_path), manifest())

    assert result.entities[0].entity_id == "e001"
    assert result.entities[0].supporting_quote == "ExampleLoader contacted 192.0.2.4."
    assert result.entities[0].mention_as_written == "ExampleLoader"


def test_import_rejects_zero_active_minutes_for_any_document(tmp_path: Path) -> None:
    path = workbook_fixture(tmp_path)
    workbook = load_workbook(path)
    for row in range(2, workbook["DOCUMENTOS"].max_row + 1):
        workbook["DOCUMENTOS"].cell(row, 10, 30)
    workbook["DOCUMENTOS"].cell(2, 10, 0)
    workbook.save(path)

    with pytest.raises(AnnotationError, match="positive integer"):
        import_annotations(path, manifest())


def test_relation_endpoint_must_exist_in_same_document(tmp_path: Path) -> None:
    with pytest.raises(AnnotationError, match="endpoint not found in document"):
        import_annotations(workbook_with_cross_doc_relation(tmp_path), manifest())


def test_import_binds_the_validated_snapshot_when_source_is_replaced(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = workbook_fixture(tmp_path / "source")
    original_hash = sha256_file(source)
    replacement = workbook_fixture(tmp_path / "replacement")
    workbook = load_workbook(replacement)
    workbook["ENTIDADES"]["G2"] = "Replacement quote."
    workbook.save(replacement)
    original_validate = annotations.validate_workbook

    def validate_then_replace(snapshot: Path, supplied_manifest: object):
        result = original_validate(snapshot, supplied_manifest)
        os.replace(replacement, source)
        return result

    monkeypatch.setattr(annotations, "validate_workbook", validate_then_replace)

    result = annotations.import_annotations(source, manifest())

    assert result.workbook_sha256 == original_hash
    assert result.entities[0].supporting_quote == "ExampleLoader contacted 192.0.2.4."
