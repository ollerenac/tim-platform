"""Behavioral contracts for the two-document blind practice package."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

import exp02.practice as practice
from exp02.practice import (
    PracticeError,
    build_practice_workbook,
    read_answer_key,
    validate_practice_workbook,
)
from exp02.workbook import SHEET_NAMES


EXP02_ROOT = Path(__file__).resolve().parents[1]
PRACTICE_ROOT = EXP02_ROOT / "practice"


def _has_annotation_rows(sheet) -> bool:  # type: ignore[no-untyped-def]
    return any(
        any(cell.value not in (None, "") for cell in row)
        for row in sheet.iter_rows(min_row=2)
    )


def test_practice_uses_two_committed_excluded_inputs_with_bound_hashes() -> None:
    """Catches practice inputs that are not the two declared excluded historical documents."""
    manifest = json.loads((PRACTICE_ROOT / "manifest.v1.json").read_text(encoding="utf-8"))
    exclusions = json.loads((EXP02_ROOT / "config" / "exclusions.v1.json").read_text(encoding="utf-8"))
    excluded = {entry["stable_id"]: entry for entry in exclusions["entries"]}

    assert manifest["schema_version"] == 1
    assert len(manifest["documents"]) == 2
    assert {document["practice_kind"] for document in manifest["documents"]} == {
        "narrative", "indicator-heavy"
    }
    for document in manifest["documents"]:
        assert document["role"] == "practice"
        historical = excluded[document["historical_stable_id"]]
        assert historical["reason"] in {"previously_evaluated", "development_material"}
        assert document["original_sha256"] == historical["sha256"]
        original = PRACTICE_ROOT / document["original_path"]
        converted = PRACTICE_ROOT / document["text_path"]
        assert hashlib.sha256(original.read_bytes()).hexdigest() == document["original_sha256"]
        assert hashlib.sha256(converted.read_bytes()).hexdigest() == document["text_sha256"]


def test_delivered_practice_workbook_is_blank_and_uses_standard_six_sheets(tmp_path: Path) -> None:
    """Catches a delivered exercise that exposes answer rows or a nonstandard layout."""
    output = tmp_path / "practice.xlsx"
    build_practice_workbook(PRACTICE_ROOT, output)

    workbook = load_workbook(output, data_only=False, keep_links=False)
    assert workbook.sheetnames == SHEET_NAMES
    assert workbook["_LISTAS"].sheet_state == "hidden"
    assert workbook["LEEME"]["B2"].value == "practice"
    assert workbook["LEEME"]["B4"].value == "two-excluded-documents"
    assert [row[0].value for row in workbook["DOCUMENTOS"].iter_rows(min_row=2)] == [
        "practice-aa25-239a-narrative",
        "practice-aa26-204a-indicators",
    ]
    assert not _has_annotation_rows(workbook["ENTIDADES"])
    assert not _has_annotation_rows(workbook["RELACIONES"])


def test_practice_reuses_annotation_layout_and_validates_saved_answers(tmp_path: Path) -> None:
    output = tmp_path / "practice.xlsx"
    build_practice_workbook(PRACTICE_ROOT, output)
    workbook = load_workbook(output)
    assert [cell.value for cell in workbook["ENTIDADES"][1]] == [
        "document_id", "entity_id", "entity_type", "indicator_subtype",
        "normalized_value", "mention_as_written", "supporting_quote",
        "page_or_lines", "certainty", "notes",
    ]
    assert workbook["DOCUMENTOS"].column_dimensions["C"].width == 40
    assert workbook["ENTIDADES"].column_dimensions["G"].width == 56
    assert workbook["RELACIONES"].column_dimensions["F"].width == 56
    assert workbook["DUDAS"].column_dimensions["C"].width == 60
    for name in ("DOCUMENTOS", "ENTIDADES", "RELACIONES", "DUDAS"):
        assert workbook[name].sheet_format.defaultRowHeight >= 30

    document_id = "practice-aa25-239a-narrative"
    workbook["DOCUMENTOS"]["J2"] = 7
    for column, value in enumerate(
        [
            document_id, "entity-smoke", "malware", None, "ExampleLoader",
            "ExampleLoader", "ExampleLoader was observed.", "lines 1-1", "clear",
            "LibreOffice smoke",
        ],
        start=1,
    ):
        workbook["ENTIDADES"].cell(2, column, value)
    workbook["DUDAS"].append(
        [document_id, "lines 1-1", "Dummy smoke question", "temporary copy"]
    )
    workbook.save(output)

    summary = validate_practice_workbook(output, PRACTICE_ROOT)

    assert summary.document_count == 2
    assert summary.entity_count == 1
    assert summary.question_count == 1


def test_committed_practice_workbook_is_sealed_blank_and_valid() -> None:
    seal = json.loads((PRACTICE_ROOT / "package.v1.json").read_text(encoding="utf-8"))
    workbook_path = PRACTICE_ROOT / "practice.xlsx"

    assert seal["schema_version"] == 1
    assert seal["practice_workbook_sha256"] == hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    assert seal["practice_manifest_sha256"] == hashlib.sha256(
        (PRACTICE_ROOT / "manifest.v1.json").read_bytes()
    ).hexdigest()
    summary = validate_practice_workbook(workbook_path, PRACTICE_ROOT)
    assert summary.document_count == 2
    assert summary.entity_count == summary.relationship_count == summary.question_count == 0


def test_practice_workbook_write_is_exclusive_and_rejects_symlink_targets(tmp_path: Path) -> None:
    """Catches a builder that overwrites a file or follows either symlink form."""
    existing = tmp_path / "practice.xlsx"
    existing.write_bytes(b"existing workbook")
    with pytest.raises(PracticeError, match="overwrite"):
        build_practice_workbook(PRACTICE_ROOT, existing)
    assert existing.read_bytes() == b"existing workbook"

    outside = tmp_path / "outside"
    outside.mkdir()
    final_link = tmp_path / "final-link" / "practice.xlsx"
    final_link.parent.mkdir()
    final_link.symlink_to(outside / "captured.xlsx")
    with pytest.raises(PracticeError, match="symlink"):
        build_practice_workbook(PRACTICE_ROOT, final_link)
    assert not (outside / "captured.xlsx").exists()

    linked_parent = tmp_path / "linked-parent"
    linked_parent.symlink_to(outside, target_is_directory=True)
    with pytest.raises(PracticeError, match="symlink"):
        build_practice_workbook(PRACTICE_ROOT, linked_parent / "practice.xlsx")
    assert not (outside / "practice.xlsx").exists()


def test_practice_workbook_rejects_a_file_created_during_serialization(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Catches a check-then-write race that could replace another writer's workbook."""
    output = tmp_path / "practice.xlsx"
    real_open_new = practice._open_new_file

    def contested_open(path: Path) -> int:
        output.write_bytes(b"other writer")
        return real_open_new(path)

    monkeypatch.setattr(practice, "_open_new_file", contested_open)
    with pytest.raises(PracticeError, match="overwrite"):
        build_practice_workbook(PRACTICE_ROOT, output)
    assert output.read_bytes() == b"other writer"


def test_answer_key_rows_use_closed_vocabulary_and_local_foreign_keys() -> None:
    """Catches an answer key with an unknown type or a relation endpoint outside its document."""
    answer = read_answer_key(EXP02_ROOT / "PRACTICE-ANSWER-KEY.md")
    manifest = json.loads((PRACTICE_ROOT / "manifest.v1.json").read_text(encoding="utf-8"))
    vocabulary = json.loads((EXP02_ROOT / "config" / "annotation-types.v1.json").read_text(encoding="utf-8"))
    document_ids = {document["document_id"] for document in manifest["documents"]}
    entities = {(row["document_id"], row["entity_id"]) for row in answer.entities}

    assert answer.entities
    assert answer.relationships
    for row in answer.entities:
        assert row["document_id"] in document_ids
        assert row["entity_type"] in vocabulary["entity_types"]
        if row["entity_type"] == "indicator":
            assert row["indicator_subtype"] in vocabulary["indicator_subtypes"]
        else:
            assert row["indicator_subtype"] is None
    for row in answer.relationships:
        assert row["document_id"] in document_ids
        assert row["relationship_type"] in vocabulary["relationship_types"]
        assert (row["document_id"], row["source_entity_id"]) in entities
        assert (row["document_id"], row["target_entity_id"]) in entities

    indicators = [
        row for row in answer.relationships
        if row["document_id"] == "practice-aa26-204a-indicators"
        and row["relationship_type"] == "indicates"
    ]
    assert len(indicators) == 2
    for row in indicators:
        assert row["page_or_lines"] == "pages 20-21"
        assert "The following indicators have been attributed to use by LAUNDRY BEAR" in (row["supporting_quote"] or "")
        assert "emailanalytics.com[.]ua 185.86.79[.]95" in (row["supporting_quote"] or "")
