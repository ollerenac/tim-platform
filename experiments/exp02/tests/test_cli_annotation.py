"""Black-box contracts for the EXP-02 annotation workflow CLI."""

from __future__ import annotations

import os
import base64
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys

import pytest
from openpyxl import load_workbook


def _canonical_json_bytes(value: object) -> bytes:
    return (
        json.dumps(
            value,
            ensure_ascii=False,
            allow_nan=False,
            sort_keys=True,
            separators=(",", ":"),
        )
        + "\n"
    ).encode("utf-8")


@pytest.fixture
def cli() -> object:
    def run(*arguments: str, cwd: Path | None = None) -> subprocess.CompletedProcess[str]:
        environment = dict(os.environ)
        environment["PYTHONPATH"] = os.pathsep.join([
            str(Path(__file__).parents[1] / "src"),
            str(Path(__file__).parents[3] / "services" / "intel-extractor"),
        ])
        return subprocess.run(
            [sys.executable, "-m", "exp02.cli", *arguments],
            text=True,
            capture_output=True,
            check=False,
            cwd=cwd,
            env=environment,
        )

    return run


@pytest.fixture
def frozen_inputs(tmp_path: Path, cli: object) -> Path:
    root = tmp_path / "evidence"
    fixture = _offline_fixture(tmp_path / "acquisition-fixture.json")
    result = _inspect_and_acquire(cli, root, fixture)
    assert result.returncode == 0, result.stderr
    return root


def test_cannot_build_workbooks_before_input_manifest(cli: object, tmp_path: Path) -> None:
    result = cli("workbooks", "build", "--root", str(tmp_path))  # type: ignore[operator]

    assert result.returncode == 2
    assert "input manifest is not frozen" in result.stderr


def test_input_verifier_rejects_legacy_extension_decision(
    cli: object, frozen_inputs: Path
) -> None:
    (frozen_inputs / "extension-decision.v1.json").write_text("{}", encoding="utf-8")

    result = cli("inputs", "verify", "--root", str(frozen_inputs))  # type: ignore[operator]

    assert result.returncode == 2
    assert "unexpected active-state artifact" in result.stderr


def test_cannot_freeze_reference_before_two_imports_and_adjudication(
    cli: object, frozen_inputs: Path
) -> None:
    result = cli("reference", "freeze", "--root", str(frozen_inputs))  # type: ignore[operator]

    assert result.returncode == 2
    assert "two annotations and adjudication are required" in result.stderr


@pytest.mark.parametrize("failure", ["candidate-fetch", "feed-fetch", "feed-parse"])
def test_external_acquisition_shortfall_has_distinct_exit_code(
    cli: object, tmp_path: Path, failure: str
) -> None:
    fixture = _offline_fixture(tmp_path / "broken-acquisition-fixture.json")
    payload = json.loads(fixture.read_text(encoding="utf-8"))
    if failure == "candidate-fetch":
        for number in range(4, 7):
            del payload["responses"][f"https://ncsc-uk.fixture.test/article-{number}"]
    elif failure == "feed-fetch":
        del payload["responses"][
            "https://www.ncsc.gov.uk/api/1/services/v1/all-rss-feed.xml"
        ]
    else:
        payload["responses"][
            "https://www.ncsc.gov.uk/api/1/services/v1/all-rss-feed.xml"
        ]["body_base64"] = base64.b64encode(b"not an RSS feed").decode()
    fixture.write_text(json.dumps(payload), encoding="utf-8")

    root = tmp_path / "evidence"
    inspected = cli("inputs", "inspect", "--root", str(root), "--fixture", str(fixture))  # type: ignore[operator]
    assert inspected.returncode == 0, inspected.stderr
    result = cli("inputs", "acquire", "--root", str(root), "--decisions", str(_decision_file(root)))  # type: ignore[operator]

    assert result.returncode == 3, (failure, result.stdout, result.stderr)
    assert "external acquisition failed" in result.stderr


def test_cannot_build_workbook_when_frozen_input_bytes_no_longer_match(
    cli: object, frozen_inputs: Path
) -> None:
    text_path = next((frozen_inputs / "inputs").glob("*/converted.txt"))
    original = text_path.read_bytes()
    text_path.write_bytes(bytes([original[0] ^ 1]) + original[1:])

    result = cli("workbooks", "build", "--root", str(frozen_inputs), "--annotator", "annotator-a")  # type: ignore[operator]

    assert result.returncode == 2
    assert "text digest mismatch" in result.stderr


def test_sources_inspect_normalizes_malformed_yaml_to_invalid_input(
    cli: object, tmp_path: Path
) -> None:
    sources = tmp_path / "sources.yaml"
    sources.write_text("sources: [unterminated", encoding="utf-8")

    result = cli("sources", "inspect", "--sources", str(sources))  # type: ignore[operator]

    assert result.returncode == 2
    assert "source registry is unavailable or invalid" in result.stderr


def test_inputs_inspect_fixture_cannot_omit_an_earlier_frozen_source(
    cli: object, tmp_path: Path
) -> None:
    fixture = _offline_fixture(tmp_path / "incomplete-registry.json")
    payload = json.loads(fixture.read_text(encoding="utf-8"))
    payload["sources"] = payload["sources"][1:]
    fixture.write_bytes(_canonical_json_bytes(payload))
    root = tmp_path / "evidence"

    result = cli(
        "inputs", "inspect", "--root", str(root), "--fixture", str(fixture)
    )  # type: ignore[operator]

    assert result.returncode == 2
    assert "input inspection is invalid" in result.stderr
    assert not root.exists()


def test_symlinked_root_is_rejected_before_any_target_write(
    cli: object, tmp_path: Path
) -> None:
    target = tmp_path / "target"
    target.mkdir()
    root = tmp_path / "linked-root"
    root.symlink_to(target, target_is_directory=True)
    fixture = _offline_fixture(tmp_path / "acquisition-fixture.json")

    result = cli("inputs", "inspect", "--root", str(root), "--fixture", str(fixture))  # type: ignore[operator]

    assert result.returncode == 2
    assert "evidence root is invalid" in result.stderr
    assert list(target.iterdir()) == []


def _offline_fixture(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    source_ids = [
        "ncsc-uk",
        "cert-eu",
        "cert-pl-en",
        "acsc-advisories",
        "unit-42",
        "eset-welivesecurity",
        "volexity",
    ]
    available_ids = {"ncsc-uk", "cert-eu", "unit-42", "eset-welivesecurity"}
    source_names = {
        "ncsc-uk": "NCSC UK",
        "cert-eu": "CERT-EU Threat Intelligence",
        "cert-pl-en": "CERT-PL EN",
        "acsc-advisories": "ACSC Advisories",
        "unit-42": "Unit 42",
        "eset-welivesecurity": "ESET WeLiveSecurity",
        "volexity": "Volexity",
    }
    feed_urls = {
        "ncsc-uk": "https://www.ncsc.gov.uk/api/1/services/v1/all-rss-feed.xml",
        "cert-eu": "https://cert.europa.eu/publications/threat-intelligence-rss",
        "cert-pl-en": "https://cert.pl/en/rss.xml",
        "acsc-advisories": "https://www.cyber.gov.au/rss/advisories",
        "unit-42": "https://unit42.paloaltonetworks.com/feed/",
        "eset-welivesecurity": "https://www.welivesecurity.com/en/rss/feed/",
        "volexity": "https://www.volexity.com/feed/",
    }
    responses: dict[str, dict[str, str]] = {}
    sources: list[dict[str, str]] = []
    for source_id in source_ids:
        feed_url = feed_urls[source_id]
        sources.append({
            "source_id": source_id,
            "source_name": source_names[source_id],
            "source_class": "institutional" if source_id in {
                "ncsc-uk", "cert-eu", "cert-pl-en", "acsc-advisories"
            } else "technical-research",
            "feed_url": feed_url,
            "publisher": source_names[source_id],
        })
        if source_id not in available_ids:
            continue
        entries = []
        for number in range(1, 7):
            article_url = f"https://{source_id}.fixture.test/article-{number}"
            entries.append(
                f"<item><title>{source_id} incident {number}</title><link>{article_url}</link>"
                f"<author>Analyst</author><pubDate>{number:02d} Aug 2026 12:00:00 +0000</pubDate></item>"
            )
            text = " ".join(["ExampleLoader contacted 192.0.2.4."] + [f"evidence{item}" for item in range(500)])
            body = (
                f"<html><body><article><h1>{source_id} report {number}</h1>"
                f"<h2>Incident narrative</h2><p>{text}</p></article></body></html>"
            )
            responses[article_url] = {"body_base64": base64.b64encode(body.encode()).decode(), "content_type": "text/html"}
        feed = "<?xml version='1.0'?><rss version='2.0'><channel>" + "".join(entries) + "</channel></rss>"
        responses[feed_url] = {"body_base64": base64.b64encode(feed.encode()).decode(), "content_type": "application/rss+xml"}
    path.write_text(json.dumps({
        "schema_version": 1,
        "sources": sources,
        "responses": responses,
    }), encoding="utf-8")
    return path


def _decision_file(root: Path) -> Path:
    inspection_path = root / "inspection-manifest.v1.json"
    inspection = json.loads(inspection_path.read_text(encoding="utf-8"))
    decisions = [
        {
            "origin_url": candidate["origin_url"],
            "original_sha256": candidate["original_sha256"],
            "text_sha256": candidate["text_sha256"],
            "threat_focused": True,
            "has_narrative_section": True,
            "language": "en",
            "translation_duplicate": False,
        }
        for candidate in inspection["candidates"]
        if "snapshot_original_path" in candidate
    ]
    path = root / "inspector-decisions.v1.json"
    path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "inspection_manifest_digest": hashlib.sha256(
                    inspection_path.read_bytes()
                ).hexdigest(),
                "decisions": decisions,
            }
        ),
        encoding="utf-8",
    )
    return path


def _inspect_and_acquire(cli: object, root: Path, fixture: Path, *, cwd: Path | None = None):  # type: ignore[no-untyped-def]
    inspected = cli(
        "inputs", "inspect", "--root", str(root), "--fixture", str(fixture), cwd=cwd
    )  # type: ignore[operator]
    assert inspected.returncode == 0, inspected.stderr
    assert (root / "inspection-manifest.v1.json").is_file()
    assert not (root / "input-manifest.v1.json").exists()
    return cli(
        "inputs", "acquire", "--root", str(root),
        "--decisions", str(_decision_file(root)), cwd=cwd,
    )  # type: ignore[operator]


def _known_answer(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    document_id = workbook["DOCUMENTOS"]["A2"].value
    for row in range(2, workbook["DOCUMENTOS"].max_row + 1):
        workbook["DOCUMENTOS"].cell(row, 10, 40)
    workbook["ENTIDADES"].append([
        document_id, "e001", "malware", None, "ExampleLoader", "ExampleLoader",
        "ExampleLoader contacted 192.0.2.4.", "lines 1-1", "clear", None,
    ])
    workbook.save(workbook_path)


def _prepared_reference(
    cli: object,
    tmp_path: Path,
    adjudication: Path | None = None,
) -> Path:
    root = tmp_path / "evidence"
    fixture = _offline_fixture(tmp_path / "acquisition-fixture.json")
    acquired = _inspect_and_acquire(cli, root, fixture)
    assert acquired.returncode == 0, acquired.stderr
    for annotator in ("annotator-a", "annotator-b"):
        result = cli("workbooks", "build", "--root", str(root), "--annotator", annotator)  # type: ignore[operator]
        assert result.returncode == 0, result.stderr
        workbook = root / "workbooks" / f"{annotator}.xlsx"
        _known_answer(workbook)
        imported = cli("annotations", "import", "--root", str(root), "--workbook", str(workbook))  # type: ignore[operator]
        assert imported.returncode == 0, imported.stderr
    compare_arguments = ["annotations", "compare", "--root", str(root)]
    if adjudication:
        compare_arguments.extend(["--adjudication", str(adjudication)])
    compared = cli(*compare_arguments)  # type: ignore[operator]
    assert compared.returncode == 0, compared.stderr
    return root


def _frozen_package(cli: object, tmp_path: Path) -> Path:
    root = _prepared_reference(cli, tmp_path)
    frozen = cli("reference", "freeze", "--root", str(root))  # type: ignore[operator]
    assert frozen.returncode == 0, frozen.stderr
    return root


def test_package_verifier_requires_complete_exact_receipt_and_full_freeze(
    cli: object, tmp_path: Path
) -> None:
    prepared = _prepared_reference(cli, tmp_path / "prepared")
    before_freeze = cli("verify-annotation-package", "--root", str(prepared))  # type: ignore[operator]
    assert before_freeze.returncode == 2

    frozen = _frozen_package(cli, tmp_path / "frozen")
    baseline = json.loads((frozen / "annotation-package.v1.json").read_text(encoding="utf-8"))
    assert cli("verify-annotation-package", "--root", str(frozen)).returncode == 0  # type: ignore[operator]

    for name, edit in (
        ("empty", lambda receipt: receipt.update({"files": {}})),
        ("missing", lambda receipt: receipt["files"].pop("annotations/reference.v1.json")),
        ("extra", lambda receipt: receipt["files"].update({"input-manifest-copy.v1.json": receipt["files"]["input-manifest.v1.json"]})),
    ):
        candidate = tmp_path / name
        shutil.copytree(frozen, candidate)
        receipt = json.loads(json.dumps(baseline))
        if name == "extra":
            shutil.copyfile(candidate / "input-manifest.v1.json", candidate / "input-manifest-copy.v1.json")
        edit(receipt)
        (candidate / "annotation-package.v1.json").write_text(json.dumps(receipt), encoding="utf-8")
        result = cli("verify-annotation-package", "--root", str(candidate))  # type: ignore[operator]
        assert result.returncode == 2, (name, result.stdout, result.stderr)

    altered = tmp_path / "altered"
    shutil.copytree(frozen, altered)
    artifact = altered / "annotations" / "reference.v1.json"
    artifact.write_bytes(artifact.read_bytes() + b" ")
    result = cli("verify-annotation-package", "--root", str(altered))  # type: ignore[operator]
    assert result.returncode == 2

    extra_artifact = tmp_path / "extra-artifact"
    shutil.copytree(frozen, extra_artifact)
    (extra_artifact / "annotations" / "unexpected.v1.json").write_text("{}", encoding="utf-8")
    result = cli("verify-annotation-package", "--root", str(extra_artifact))  # type: ignore[operator]
    assert result.returncode == 2


def test_reference_freeze_preflight_leaves_no_partial_artifacts(
    cli: object, tmp_path: Path
) -> None:
    receipt_root = _prepared_reference(cli, tmp_path / "receipt")
    (receipt_root / "annotation-package.v1.json").write_text("{}", encoding="utf-8")
    receipt_failure = cli("reference", "freeze", "--root", str(receipt_root))  # type: ignore[operator]
    assert receipt_failure.returncode == 2
    assert not (receipt_root / "annotations").exists()

    outside = tmp_path / "outside.xlsx"
    outside_root = _prepared_reference(cli, tmp_path / "outside-root")
    outside_failure = cli("reference", "freeze", "--root", str(outside_root), "--adjudication", str(outside))  # type: ignore[operator]
    assert outside_failure.returncode == 2
    assert not (outside_root / "annotations").exists()


def test_offline_cli_workflow_freezes_and_detects_converted_text_mutation(
    cli: object, tmp_path: Path
) -> None:
    root = _frozen_package(cli, tmp_path)
    receipt = json.loads((root / "annotation-package.v1.json").read_text(encoding="utf-8"))
    assert "extension-decision.v1.json" not in receipt["files"]
    assert "inspection-manifest.v1.json" in receipt["files"]
    assert "inspector-decisions.v1.json" in receipt["files"]
    assert "workbooks/annotator-a.xlsx" in receipt["files"]
    assert "workbooks/annotator-b.xlsx" in receipt["files"]
    assert cli("verify-annotation-package", "--root", str(root)).returncode == 0  # type: ignore[operator]

    text_path = next((root / "inputs").glob("*/converted.txt"))
    original = text_path.read_bytes()
    text_path.write_bytes(bytes([original[0] ^ 1]) + original[1:])
    mutated = cli("verify-annotation-package", "--root", str(root))  # type: ignore[operator]

    assert mutated.returncode == 2
    assert "text digest mismatch" in mutated.stderr

    text_path.write_bytes(original)
    inspection_path = next((root / "inspection").glob("*/converted.txt"))
    inspected = inspection_path.read_bytes()
    inspection_path.write_bytes(bytes([inspected[0] ^ 1]) + inspected[1:])
    mutated_inspection = cli(
        "verify-annotation-package", "--root", str(root)
    )  # type: ignore[operator]

    assert mutated_inspection.returncode == 2
    assert "inspection snapshot digest mismatch" in mutated_inspection.stderr


def test_extension_command_and_24_document_workbook_are_not_available(
    cli: object, frozen_inputs: Path
) -> None:
    extension = cli("extension", "decide", "--root", str(frozen_inputs))  # type: ignore[operator]
    workbook = cli(
        "workbooks", "build", "--root", str(frozen_inputs),
        "--annotator", "annotator-a",
        "--output", str(frozen_inputs / "workbooks" / "annotator-a-24.xlsx"),
    )  # type: ignore[operator]

    assert extension.returncode == 2
    assert workbook.returncode == 2


def test_input_verifier_rejects_self_consistent_undersized_manifest(
    cli: object, frozen_inputs: Path, tmp_path: Path
) -> None:
    baseline = json.loads(
        (frozen_inputs / "input-manifest.v1.json").read_text(encoding="utf-8")
    )

    def undersized(manifest: dict[str, object]) -> None:
        manifest["documents"] = manifest["documents"][:-1]  # type: ignore[index]

    def missing_source(manifest: dict[str, object]) -> None:
        manifest["sources"] = manifest["sources"][:-1]  # type: ignore[index]

    def reordered_sources(manifest: dict[str, object]) -> None:
        manifest["sources"][0], manifest["sources"][1] = (  # type: ignore[index]
            manifest["sources"][1], manifest["sources"][0]  # type: ignore[index]
        )

    def wrong_role(manifest: dict[str, object]) -> None:
        manifest["documents"][0]["role"] = "extension"  # type: ignore[index]

    def legacy_twenty_four(manifest: dict[str, object]) -> None:
        documents = manifest["documents"]  # type: ignore[assignment]
        template = documents[0]
        for number in range(17, 25):
            document_id = f"legacy-{number}"
            documents.append(
                {
                    **template,
                    "document_id": document_id,
                    "original_path": f"experiments/exp02/evidence/inputs/{document_id}/original.html",
                    "text_path": f"experiments/exp02/evidence/inputs/{document_id}/converted.txt",
                }
            )

    def wrong_policy(manifest: dict[str, object]) -> None:
        manifest["policy_digest"] = "0" * 64

    def wrong_membership(manifest: dict[str, object]) -> None:
        manifest["documents"][0]["origin_url"] = "https://example.test/not-selected"  # type: ignore[index]

    for name, mutate in (
        ("undersized", undersized),
        ("missing-source", missing_source),
        ("reordered", reordered_sources),
        ("wrong-role", wrong_role),
        ("legacy-24", legacy_twenty_four),
        ("wrong-policy", wrong_policy),
        ("wrong-membership", wrong_membership),
    ):
        candidate = tmp_path / name
        shutil.copytree(frozen_inputs, candidate)
        manifest = json.loads(json.dumps(baseline))
        mutate(manifest)
        (candidate / "input-manifest.v1.json").write_text(
            json.dumps(manifest), encoding="utf-8"
        )

        result = cli("inputs", "verify", "--root", str(candidate))  # type: ignore[operator]

        assert result.returncode == 2, name
        assert "input manifest" in result.stderr


def _rewrite_final_evidence_away_from_inspection(root: Path) -> None:
    manifest_path = root / "input-manifest.v1.json"
    ledger_path = root / "candidate-ledger.v1.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    ledger = json.loads(ledger_path.read_text(encoding="utf-8"))
    record = manifest["documents"][0]
    original_path = (
        root
        / "inputs"
        / record["document_id"]
        / f"original.{record['media_type']}"
    )
    text_path = root / "inputs" / record["document_id"] / "converted.txt"
    original_path.write_bytes(original_path.read_bytes() + b"rewritten")
    text_path.write_bytes(text_path.read_bytes() + b" rewritten")
    rewritten_original = hashlib.sha256(original_path.read_bytes()).hexdigest()
    rewritten_text = hashlib.sha256(text_path.read_bytes()).hexdigest()
    record["original_sha256"] = rewritten_original
    record["text_sha256"] = rewritten_text
    candidate = next(
        item
        for item in ledger["candidates"]
        if item.get("origin_url") == record["origin_url"]
    )
    candidate["original_sha256"] = rewritten_original
    candidate["text_sha256"] = rewritten_text
    ledger_path.write_bytes(_canonical_json_bytes(ledger))
    manifest["candidate_ledger_digest"] = hashlib.sha256(
        ledger_path.read_bytes()
    ).hexdigest()
    manifest_path.write_bytes(_canonical_json_bytes(manifest))


def test_input_verifier_rejects_recanonicalized_final_evidence_that_diverges_from_inspection(
    cli: object, frozen_inputs: Path
) -> None:
    _rewrite_final_evidence_away_from_inspection(frozen_inputs)

    result = cli("inputs", "verify", "--root", str(frozen_inputs))  # type: ignore[operator]

    assert result.returncode == 2
    assert "inspection provenance" in result.stderr


@pytest.mark.parametrize(
    ("field", "replacement"),
    [
        ("author", "Different Organization"),
        ("author_basis", "source_publisher"),
    ],
)
def test_input_verifier_rejects_altered_author_provenance(
    cli: object,
    frozen_inputs: Path,
    field: str,
    replacement: str,
) -> None:
    """Changing either resolved author field must break inspected provenance."""
    manifest_path = frozen_inputs / "input-manifest.v1.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["documents"][0]["author_basis"] == "feed_author"
    manifest["documents"][0][field] = replacement
    manifest_path.write_bytes(_canonical_json_bytes(manifest))

    result = cli("inputs", "verify", "--root", str(frozen_inputs))  # type: ignore[operator]

    assert result.returncode == 2
    assert "input manifest" in result.stderr


def test_package_verifier_rejects_recanonicalized_final_evidence_that_diverges_from_inspection(
    cli: object, tmp_path: Path
) -> None:
    root = _frozen_package(cli, tmp_path)
    _rewrite_final_evidence_away_from_inspection(root)

    result = cli("verify-annotation-package", "--root", str(root))  # type: ignore[operator]

    assert result.returncode == 2
    assert "inspection provenance" in result.stderr


@pytest.mark.parametrize(
    ("field", "replacement"),
    [
        ("author", "Different Organization"),
        ("author_basis", "source_publisher"),
    ],
)
def test_package_verifier_rejects_altered_author_provenance(
    cli: object,
    tmp_path: Path,
    field: str,
    replacement: str,
) -> None:
    root = _frozen_package(cli, tmp_path)
    manifest_path = root / "input-manifest.v1.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["documents"][0]["author_basis"] == "feed_author"
    manifest["documents"][0][field] = replacement
    manifest_path.write_bytes(_canonical_json_bytes(manifest))

    result = cli("verify-annotation-package", "--root", str(root))  # type: ignore[operator]

    assert result.returncode == 2


def test_input_verifier_rejects_a_rewritten_closed_inspector_decision_file(
    cli: object, frozen_inputs: Path
) -> None:
    decisions_path = frozen_inputs / "inspector-decisions.v1.json"
    decisions = json.loads(decisions_path.read_text(encoding="utf-8"))
    decisions["decisions"][0]["language"] = "es"
    decisions_path.write_bytes(_canonical_json_bytes(decisions))

    result = cli("inputs", "verify", "--root", str(frozen_inputs))  # type: ignore[operator]

    assert result.returncode == 2
    assert "inspector decisions digest mismatch" in result.stderr


def test_bound_workbook_is_reimported_before_comparison(
    cli: object, frozen_inputs: Path
) -> None:
    for annotator in ("annotator-a", "annotator-b"):
        assert cli("workbooks", "build", "--root", str(frozen_inputs), "--annotator", annotator).returncode == 0  # type: ignore[operator]
        workbook = frozen_inputs / "workbooks" / f"{annotator}.xlsx"
        _known_answer(workbook)
        assert cli("annotations", "import", "--root", str(frozen_inputs), "--workbook", str(workbook)).returncode == 0  # type: ignore[operator]
    workbook = frozen_inputs / "workbooks" / "annotator-a.xlsx"
    changed = load_workbook(workbook)
    changed["ENTIDADES"]["G2"] = "Changed after import."
    changed.save(workbook)

    result = cli("annotations", "compare", "--root", str(frozen_inputs))  # type: ignore[operator]

    assert result.returncode == 2


def test_documented_relative_root_completes_freeze_and_verification(
    cli: object, tmp_path: Path
) -> None:
    cwd = tmp_path / "documented-cwd"
    cwd.mkdir()
    root_argument = Path("experiments/exp02/evidence")
    root = cwd / root_argument
    fixture = _offline_fixture(cwd / "acquisition-fixture.json")

    inspected = cli(
        "inputs", "inspect", "--root", str(root_argument),
        "--fixture", str(fixture), cwd=cwd,
    )  # type: ignore[operator]
    assert inspected.returncode == 0, inspected.stderr
    decisions = _decision_file(root)
    acquired = cli(
        "inputs", "acquire", "--root", str(root_argument),
        "--decisions", str(decisions), cwd=cwd,
    )  # type: ignore[operator]
    assert acquired.returncode == 0, acquired.stderr
    for annotator in ("annotator-a", "annotator-b"):
        workbook_argument = root_argument / "workbooks" / f"{annotator}.xlsx"
        built = cli("workbooks", "build", "--root", str(root_argument), "--annotator", annotator, cwd=cwd)  # type: ignore[operator]
        assert built.returncode == 0, built.stderr
        _known_answer(root / "workbooks" / f"{annotator}.xlsx")
        imported = cli("annotations", "import", "--root", str(root_argument), "--workbook", str(workbook_argument), cwd=cwd)  # type: ignore[operator]
        assert imported.returncode == 0, imported.stderr
    assert cli("annotations", "compare", "--root", str(root_argument), cwd=cwd).returncode == 0  # type: ignore[operator]
    frozen = cli("reference", "freeze", "--root", str(root_argument), cwd=cwd)  # type: ignore[operator]
    assert frozen.returncode == 0, frozen.stderr
    verified = cli("verify-annotation-package", "--root", str(root_argument), cwd=cwd)  # type: ignore[operator]
    assert verified.returncode == 0, verified.stderr
