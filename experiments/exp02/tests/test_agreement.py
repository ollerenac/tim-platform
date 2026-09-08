"""Contracts for deterministic comparison and reference adjudication."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from exp02.agreement import (
    apply_adjudication,
    build_adjudication_workbook,
    compare_annotation_sets,
    freeze_reference,
    validate_adjudication_workbook,
)
from exp02.annotations import AnnotationSet, EntityAnnotation, RelationshipAnnotation


DOCUMENT_IDS = ("doc-one", "doc-two", *[f"doc-{number:02d}" for number in range(3, 17)])


def _annotation(annotator_id: str, value: str, quote: str) -> AnnotationSet:
    return AnnotationSet(
        annotator_id=annotator_id,
        manifest_digest="a" * 64,
        workbook_sha256="b" * 64,
        entities=(EntityAnnotation(
            document_id="doc-one", entity_id="e001", entity_type="malware",
            indicator_subtype=None, normalized_value=value, mention_as_written=value,
            supporting_quote=quote, page_or_lines="lines 1-1", certainty="clear", notes=None,
        ),),
        relationships=(),
        document_ids=DOCUMENT_IDS,
    )


def test_comparison_reports_quote_disagreement_without_losing_entity_match() -> None:
    comparison = compare_annotation_sets(
        _annotation("annotator-a", "ExampleLoader", "ExampleLoader contacted 192.0.2.4."),
        _annotation("annotator-b", "ExampleLoader", "ExampleLoader contacted 192.0.2.4 yesterday."),
    )

    assert comparison.entity_metrics.tp == 1
    assert comparison.entity_metrics.f1 == 1.0
    assert comparison.quote_agreement.matched_claims == 1
    assert comparison.quote_agreement.exact_matches == 0
    assert comparison.disagreements[0].claim_kind == "entity"


def _known_answer_sets() -> tuple[AnnotationSet, AnnotationSet]:
    common_a = EntityAnnotation("doc-one", "e001", "malware", None, "ExampleLoader", "ExampleLoader", "ExampleLoader contacted 192.0.2.4.", "lines 1-1", "clear", None)
    common_b = EntityAnnotation("doc-one", "x001", "malware", None, "exampleloader", "ExampleLoader", "ExampleLoader contacted 192.0.2.4.", "lines 1-1", "clear", None)
    alias_a = EntityAnnotation("doc-one", "e002", "threat-actor", None, "APT Zebra", "APT Zebra", "APT Zebra used ExampleLoader.", "lines 2-2", "clear", None)
    alias_b = EntityAnnotation("doc-one", "x002", "threat-actor", None, "Zebra Group", "Zebra Group", "Zebra Group used ExampleLoader.", "lines 2-2", "clear", None)
    quote_a = EntityAnnotation("doc-two", "e003", "indicator", "ip", "192.0.2.4", "192.0.2.4", "Observed 192.0.2.4.", "lines 4-4", "clear", None)
    quote_b = EntityAnnotation("doc-two", "x003", "indicator", "ip", "192.0.2.4", "192.0.2.4", "The report observed 192.0.2.4.", "lines 4-4", "clear", None)
    relation_a = RelationshipAnnotation("doc-one", "r001", "e002", "uses", "e001", "APT Zebra used ExampleLoader.", "lines 2-2", "clear", None)
    relation_b = RelationshipAnnotation("doc-one", "s001", "x002", "targets", "x001", "Zebra Group targets ExampleLoader.", "lines 2-2", "clear", None)
    a = AnnotationSet("annotator-a", "a" * 64, "b" * 64, (common_a, alias_a, quote_a), (relation_a,), DOCUMENT_IDS)
    b = AnnotationSet("annotator-b", "a" * 64, "c" * 64, (common_b, alias_b, quote_b), (relation_b,), DOCUMENT_IDS)
    return a, b


def test_known_answer_comparison_is_one_to_one_and_disagreements_are_stable() -> None:
    a, b = _known_answer_sets()
    first = compare_annotation_sets(a, b)
    second = compare_annotation_sets(a, b)

    assert (first.entity_metrics.tp, first.entity_metrics.fp, first.entity_metrics.fn) == (2, 1, 1)
    assert first.entity_metrics.f1 == pytest.approx(2 / 3)
    assert (first.relationship_metrics.tp, first.relationship_metrics.fp, first.relationship_metrics.fn) == (0, 1, 1)
    assert first.relationship_metrics_reversed.f1 == first.relationship_metrics.f1
    assert first.quote_agreement.exact_matches == 1
    assert [item.disagreement_id for item in first.disagreements] == [item.disagreement_id for item in second.disagreements]
    # Each unmatched side is its own adjudicable claim; the quote mismatch is a
    # single paired disagreement, so this fixture has five rows in total.
    assert len(first.disagreements) == 5


def test_adjudication_and_freeze_emit_only_bound_human_artifacts(tmp_path: Path) -> None:
    comparison = compare_annotation_sets(*_known_answer_sets())
    workbook = tmp_path / "adjudication.xlsx"
    build_adjudication_workbook(comparison, workbook)
    assert workbook.exists()

    decisions = {
        item.disagreement_id: {"decision": "accept-a" if item.annotator_a_value else "accept-b"}
        for item in comparison.disagreements
    }
    reference = apply_adjudication(comparison, decisions)
    frozen = freeze_reference(reference, tmp_path / "freeze", frozen_at="2026-08-30T12:00:00Z")

    assert len(reference.entities) == 4
    assert reference.relationships[0].source_entity_id in {entity.entity_id for entity in reference.entities}
    assert frozen.document_count == 16
    assert set(path.name for path in (tmp_path / "freeze" / "annotations").iterdir()) == {
        "annotator-a.v1.json", "annotator-b.v1.json", "agreement.v1.json", "adjudication.v1.json",
        "reference.v1.json", "reference-freeze.v1.json",
    }


def test_comparison_requires_two_exact_fixed_sixteen_coverages() -> None:
    complete = _annotation("annotator-a", "ExampleLoader", "same quote")
    incomplete = _annotation("annotator-b", "ExampleLoader", "same quote")
    incomplete = AnnotationSet(
        incomplete.annotator_id,
        incomplete.manifest_digest,
        incomplete.workbook_sha256,
        incomplete.entities,
        incomplete.relationships,
        DOCUMENT_IDS[:-1],
    )

    with pytest.raises(ValueError, match="exactly 16"):
        compare_annotation_sets(complete, incomplete)


def test_comparison_orders_canonical_identities_and_rejects_duplicates() -> None:
    a = _annotation("annotator-a", "ExampleLoader", "same quote")
    b = _annotation("annotator-b", "ExampleLoader", "same quote")

    comparison = compare_annotation_sets(b, a)

    assert comparison.annotator_a.annotator_id == "annotator-a"
    assert comparison.annotator_b.annotator_id == "annotator-b"
    with pytest.raises(ValueError, match="annotator-a and annotator-b"):
        compare_annotation_sets(a, a)


def test_indicator_comparison_normalizes_by_subtype_without_folding_url_path() -> None:
    def annotation(annotator: str, values: tuple[tuple[str, str], ...]) -> AnnotationSet:
        return AnnotationSet(
            annotator,
            "a" * 64,
            ("b" if annotator == "annotator-a" else "c") * 64,
            tuple(
                EntityAnnotation(
                    "doc-one", f"e-{index}", "indicator", subtype, value, value,
                    f"Observed {value}.", "lines 1-1", "clear", None,
                )
                for index, (subtype, value) in enumerate(values)
            ),
            (),
            DOCUMENT_IDS,
        )

    a = annotation(
        "annotator-a",
        (
            ("url", "HTTPS://Example.COM/Case?Token=AbC#Frag"),
            ("domain", "EXAMPLE.COM."),
            ("ip", "2001:0db8::1"),
            ("hash_sha256", "A" * 64),
            ("email", "Analyst@EXAMPLE.COM"),
        ),
    )
    b = annotation(
        "annotator-b",
        (
            ("url", "https://example.com/case?Token=AbC#Frag"),
            ("domain", "example.com"),
            ("ip", "2001:db8:0:0:0:0:0:1"),
            ("hash_sha256", "a" * 64),
            ("email", "Analyst@example.com"),
        ),
    )

    comparison = compare_annotation_sets(a, b)

    assert comparison.entity_metrics.tp == 4
    assert comparison.entity_metrics.fp == 1
    assert comparison.entity_metrics.fn == 1


def test_adjudication_rejects_impossible_side_choices_and_applies_relationship_endpoints() -> None:
    a_entities = (
        EntityAnnotation("doc-one", "a-source", "malware", None, "Source", "Source", "Source appeared.", "lines 1-1", "clear", None),
        EntityAnnotation("doc-one", "a-target", "organization", None, "Target", "Target", "Target appeared.", "lines 2-2", "clear", None),
    )
    b_entities = (
        EntityAnnotation("doc-one", "b-source", "malware", None, "Source", "Source", "Source appeared.", "lines 1-1", "clear", None),
        EntityAnnotation("doc-one", "b-target", "organization", None, "Target", "Target", "Target appeared.", "lines 2-2", "clear", None),
    )
    relation = RelationshipAnnotation(
        "doc-one", "r-one", "a-source", "uses", "a-target",
        "Source uses Target.", "lines 3-3", "clear", None,
    )
    comparison = compare_annotation_sets(
        AnnotationSet("annotator-a", "a" * 64, "b" * 64, a_entities, (relation,), DOCUMENT_IDS),
        AnnotationSet("annotator-b", "a" * 64, "c" * 64, b_entities, (), DOCUMENT_IDS),
    )
    disagreement = comparison.disagreements[0]

    with pytest.raises(ValueError, match="accept-b"):
        apply_adjudication(comparison, {disagreement.disagreement_id: {"decision": "accept-b"}})
    with pytest.raises(ValueError, match="accept-both"):
        apply_adjudication(comparison, {disagreement.disagreement_id: {"decision": "accept-both"}})
    with pytest.raises(ValueError, match="endpoint not found"):
        apply_adjudication(
            comparison,
            {
                disagreement.disagreement_id: {
                    "decision": "replace",
                    "final_type": "targets",
                    "final_value": "missing -> a-source",
                    "final_quote": "Missing targets Source.",
                    "rationale": "Exercise the foreign-key gate.",
                }
            },
        )

    reference = apply_adjudication(
        comparison,
        {
            disagreement.disagreement_id: {
                "decision": "replace",
                "final_type": "targets",
                "final_value": "a-target -> a-source",
                "final_quote": "Target targets Source.",
                "rationale": "The adjudicators reversed the endpoints.",
            }
        },
    )
    by_id = {entity.entity_id: entity.normalized_value for entity in reference.entities}
    result = reference.relationships[0]
    assert result.relationship_type == "targets"
    assert by_id[result.source_entity_id] == "Target"
    assert by_id[result.target_entity_id] == "Source"


def test_adjudication_workbook_uses_exclusive_safe_write_and_ooxml_preflight(
    tmp_path: Path,
) -> None:
    comparison = compare_annotation_sets(
        _annotation("annotator-a", "ExampleLoader", "first quote"),
        _annotation("annotator-b", "ExampleLoader", "second quote"),
    )
    path = tmp_path / "adjudication.xlsx"
    build_adjudication_workbook(comparison, path)
    original = path.read_bytes()
    with pytest.raises(ValueError, match="overwrite"):
        build_adjudication_workbook(comparison, path)
    assert path.read_bytes() == original

    macro = tmp_path / "macro.xlsx"
    macro.write_bytes(original)
    with ZipFile(macro, "a", ZIP_DEFLATED) as archive:
        archive.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(ValueError, match="macros"):
        validate_adjudication_workbook(macro, comparison)

    external = tmp_path / "external.xlsx"
    external.write_bytes(original)
    with ZipFile(external, "a", ZIP_DEFLATED) as archive:
        archive.writestr("xl/externalLinks/externalLink1.xml", b"<externalLink/>")
    with pytest.raises(ValueError, match="external links"):
        validate_adjudication_workbook(external, comparison)

    formula = tmp_path / "formula.xlsx"
    formula.write_bytes(original)
    workbook = load_workbook(formula)
    workbook.active["F2"] = '=HYPERLINK("https://example.test","accept-a")'
    workbook.save(formula)
    with pytest.raises(ValueError, match="formula"):
        validate_adjudication_workbook(formula, comparison)

    outside = tmp_path / "outside"
    outside.mkdir()
    linked = tmp_path / "linked"
    linked.symlink_to(outside, target_is_directory=True)
    with pytest.raises(ValueError, match="symlink"):
        build_adjudication_workbook(comparison, linked / "adjudication.xlsx")
    assert list(outside.iterdir()) == []


def test_adjudication_rejects_unknown_keys_and_model_evidence_before_freeze() -> None:
    comparison = compare_annotation_sets(
        _annotation("annotator-a", "ExampleLoader", "first quote"),
        _annotation("annotator-b", "ExampleLoader", "second quote"),
    )
    identifier = comparison.disagreements[0].disagreement_id

    with pytest.raises(ValueError, match="unknown decision field"):
        apply_adjudication(comparison, {identifier: {"decision": "accept-a", "tim_condition": "C4"}})
    with pytest.raises(ValueError, match="model-result"):
        apply_adjudication(comparison, {
            identifier: {
                "decision": "replace", "final_type": "malware", "final_value": "ExampleLoader",
                "final_quote": "reviewed quote", "rationale": "runs/c4/output.json",
            }
        })


def test_adjudication_workbook_rejects_extra_columns_and_altered_claim_displays(tmp_path: Path) -> None:
    comparison = compare_annotation_sets(*_known_answer_sets())
    extra_column = tmp_path / "extra-column.xlsx"
    build_adjudication_workbook(comparison, extra_column)
    workbook = load_workbook(extra_column)
    workbook.active.cell(1, 12, "tim_condition")
    workbook.save(extra_column)
    with pytest.raises(ValueError, match="exactly"):
        validate_adjudication_workbook(extra_column, comparison)

    altered_claim = tmp_path / "altered-claim.xlsx"
    build_adjudication_workbook(comparison, altered_claim)
    workbook = load_workbook(altered_claim)
    workbook.active["D2"] = "not the imported annotator claim"
    workbook.save(altered_claim)
    with pytest.raises(ValueError, match="immutable claim"):
        validate_adjudication_workbook(altered_claim, comparison)


def test_adjudication_workbook_returns_only_the_closed_decision_record(tmp_path: Path) -> None:
    comparison = compare_annotation_sets(*_known_answer_sets())
    path = tmp_path / "closed-decision.xlsx"
    build_adjudication_workbook(comparison, path)
    workbook = load_workbook(path)
    for row, disagreement in enumerate(comparison.disagreements, start=2):
        workbook.active.cell(
            row, 6, "accept-a" if disagreement.annotator_a_value else "accept-b"
        )
    workbook.save(path)

    decisions = validate_adjudication_workbook(path, comparison)

    assert all(set(record) == {"decision"} for record in decisions.values())


def test_freeze_rejects_non_utc_timestamps(tmp_path: Path) -> None:
    comparison = compare_annotation_sets(*_known_answer_sets())
    decisions = {
        item.disagreement_id: {"decision": "accept-a" if item.annotator_a_value else "accept-b"}
        for item in comparison.disagreements
    }
    reference = apply_adjudication(comparison, decisions)

    for name, timestamp in (("offset", "2026-08-30T12:00:00+00:00"), ("invalid", "not-a-time")):
        with pytest.raises(ValueError, match="UTC"):
            freeze_reference(reference, tmp_path / name, frozen_at=timestamp)
